"""Bank desk: statements, counterparties and their classification."""

import datetime
from decimal import Decimal
from io import BytesIO

import openpyxl
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse

from finance_saber.cashbook import bank
from finance_saber.cashbook.ledger import balances
from finance_saber.cashbook.models import Account, Article, CashDesk, Counterparty, Kind, Operation, Region

D = Decimal
SALES = "20208000900713230001/304577774/GOODWILL BUSINESS TRADE xususiy korxonasi"
SUPPLIER = "20210000900126024001/201037628/O`Zbekekspertiza AJ"
NEW = "20208000800185981007/200640719/TOSHKENT XALQARO AEROPORTI MCHJ"


@pytest.fixture
def desk(db):
    return CashDesk.objects.get(slug="bank")  # added by migration 0007


@pytest.fixture
def counterparties(db):
    sales = Article.objects.get(name="Месный продажа")
    supplies = Article.objects.get(name="Местные поставщики")
    tash = Region.objects.get(name="Tash reg")
    return {
        "sales": Counterparty.objects.create(
            reference=SALES, **Counterparty.parse_reference(SALES), article=sales, region=tash
        ),
        "supplier": Counterparty.objects.create(
            reference=SUPPLIER, **Counterparty.parse_reference(SUPPLIER), article=supplies
        ),
    }


def statement(rows, title="Выписка"):
    """An .xlsx statement like the bank's export, as bytes."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append([title])  # a title line above the header, as bank exports have
    # A date-only helper column first, then the real one; "Cчет" with a Latin C, as in the export.
    sheet.append(
        ["Дата", "Дата", "Cчет/ИНН", "№ док", "Оп", "МФО", "Оборот Дебет", "Оборот Кредит", "Назначение платежа"]
    )
    for posted_at, reference, document, debit, credit, purpose in rows:
        sheet.append(
            [posted_at.date(), posted_at, reference, document, 4 if credit else 1, "00440", debit, credit, purpose]
        )
    sheet.append([None, None, "Итого", None, None, None, 999, 999, None])  # totals line: ignored
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


AT = datetime.datetime(2026, 10, 20, 15, 47, 3)
ROWS = [
    (AT, SALES, "517", None, 72000000, "Оплата за напитки"),
    (AT, SALES, "518", 1500000, None, "Возврат за брак"),  # a refund: sales article used for money out
    (AT, SUPPLIER, "1249", 461440, None, "Оплата услуг экспертизы"),
    (AT, NEW, "77", 250000, None, "Аэропорт"),
    (AT, NEW, "77", 250000, None, "Аэропорт"),  # identical second payment: both count
]


def test_read_statement_finds_columns_by_header():
    lines = bank.read_statement(BytesIO(statement(ROWS)))
    assert len(lines) == 5
    first = lines[0]
    assert (first.posted_at, first.reference, first.document, first.credit, first.kind) == (
        AT, SALES, "517", D("72000000.00"), Kind.INCOME,
    )  # fmt: skip
    assert lines[2].debit == D("461440.00") and lines[2].kind == Kind.EXPENSE
    assert [line.occurrence for line in lines[3:]] == [1, 2]
    assert lines[3].key != lines[4].key


def test_read_statement_reads_the_bank_sheet_of_the_finance_workbook():
    workbook = openpyxl.load_workbook(BytesIO(statement(ROWS)))
    workbook.active.title = "Банк"
    workbook.create_sheet("Касса Бон", 0)  # not the first sheet, as in Situation.xlsx
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    assert len(bank.read_statement(buffer)) == 5


def test_read_statement_rejects_other_files():
    workbook = openpyxl.Workbook()
    workbook.active.append(["Something", "else"])
    buffer = BytesIO()
    workbook.save(buffer)
    with pytest.raises(bank.StatementError):
        bank.read_statement(BytesIO(buffer.getvalue()))
    with pytest.raises(bank.StatementError):
        bank.read_statement(BytesIO(b"not a workbook"))


def test_record_lines_classifies_by_counterparty(desk, counterparties):
    report = bank.record_lines(desk, bank.read_statement(BytesIO(statement(ROWS))))
    assert (report.created, report.duplicates, report.unclassified) == (5, 0, 2)
    assert [c.reference for c in report.new_counterparties] == [NEW]

    sale = desk.operations.get(document_number="517")
    assert (sale.kind, sale.category.article.name, sale.counterparty.region.name) == (
        Kind.INCOME, "Месный продажа", "Tash reg",
    )  # fmt: skip
    refund = desk.operations.get(document_number="518")
    assert (refund.kind, refund.category.article.name) == (Kind.EXPENSE, "Месный продажа")  # "both" article
    assert desk.operations.filter(document_number="77", category__name="Требует проверки").count() == 2
    assert balances(desk.operations.all())[Account.BANK] == D("72000000") - D("1500000") - D("461440") - D("500000")


def test_uploading_the_same_statement_again_adds_nothing(desk, counterparties):
    data = statement(ROWS)
    bank.record_lines(desk, bank.read_statement(BytesIO(data)))
    report = bank.record_lines(desk, bank.read_statement(BytesIO(data)))
    assert (report.created, report.duplicates) == (0, 5)
    assert desk.operations.count() == 5


def test_classifying_a_counterparty_refiles_its_lines(client, user, desk, counterparties):
    bank.record_lines(desk, bank.read_statement(BytesIO(statement(ROWS))))
    airport = Counterparty.objects.get(reference=NEW)
    other = Article.objects.get(name="Прочие расходы")
    client.force_login(user)
    response = client.post(
        reverse("cashbook:counterparty_edit", args=[airport.pk]),
        {"name": airport.name, "inn": airport.inn, "account_number": airport.account_number,
         "article": other.pk, "region": "", "apply_to_existing": "on"},
    )  # fmt: skip
    assert response.status_code == 302
    assert set(airport.operations.values_list("category__article__name", flat=True)) == {"Прочие расходы"}


def test_upload_page(client, user, desk, counterparties):
    client.force_login(user)
    upload = SimpleUploadedFile("statement.xlsx", statement(ROWS))
    response = client.post(reverse("cashbook:statement_upload", kwargs={"desk": "bank"}), {"file": upload})
    assert response.status_code == 302
    assert "date_from=2026-10-20" in response.url
    assert desk.operations.count() == 5

    page = client.get(response.url)
    content = page.content.decode()
    assert "GOODWILL BUSINESS TRADE" in content
    assert "Tash reg" in content
    assert page.context["unclassified_count"] == 1


def test_upload_rejects_wrong_files(client, user, desk):
    client.force_login(user)
    url = reverse("cashbook:statement_upload", kwargs={"desk": "bank"})
    client.post(url, {"file": SimpleUploadedFile("statement.csv", b"a,b")})
    client.post(url, {"file": SimpleUploadedFile("statement.xlsx", b"not a workbook")})
    assert not Operation.objects.exists()


def test_upload_only_on_bank_desks(client, user):
    client.force_login(user)
    url = reverse("cashbook:statement_upload", kwargs={"desk": "bon"})
    assert client.post(url, {"file": SimpleUploadedFile("s.xlsx", statement(ROWS))}).status_code == 404


def test_counterparty_list_filters_unclassified(client, user, desk, counterparties):
    bank.record_lines(desk, bank.read_statement(BytesIO(statement(ROWS))))
    client.force_login(user)
    response = client.get(reverse("cashbook:counterparties") + "?status=unclassified")
    assert [c.reference for c in response.context["counterparties"]] == [NEW]
    response = client.get(reverse("cashbook:counterparties") + "?q=304577774")
    assert [c.reference for c in response.context["counterparties"]] == [SALES]


def test_bank_entry_form_has_counterparty_and_document(client, user, desk, counterparties):
    client.force_login(user)
    sale_category = Article.objects.get(name="Месный продажа").categories.get(name="Месный продажа")
    response = client.post(
        reverse("cashbook:operations", kwargs={"desk": "bank"}),
        {"kind": "income", "date": "2026-10-20", "category": sale_category.pk, "account": "bank",
         "amount": "1 000 000", "counterparty": counterparties["sales"].pk, "document_number": "900"},
    )  # fmt: skip
    assert response.status_code == 302
    operation = desk.operations.get()
    assert (operation.counterparty, operation.document_number) == (counterparties["sales"], "900")


def test_parse_reference():
    assert Counterparty.parse_reference(SALES) == {
        "account_number": "20208000900713230001",
        "inn": "304577774",
        "name": "GOODWILL BUSINESS TRADE xususiy korxonasi",
    }
    assert Counterparty.parse_reference("Наличные")["name"] == "Наличные"


def test_toolbar_panels_open_when_needed(client, user, desk):
    client.force_login(user)
    bank_page = client.get(reverse("cashbook:operations", kwargs={"desk": "bank"}))
    assert not bank_page.context["entry_open"]  # the bank works from statements
    assert not bank_page.context["filters_open"]
    assert bank_page.context["active_filters"] == 0

    cash_page = client.get(reverse("cashbook:operations", kwargs={"desk": "bon"}))
    assert not cash_page.context["entry_open"]  # every desk: the form opens from its button
    inkassa_page = client.get(reverse("cashbook:operations", kwargs={"desk": "inkassa"}))
    assert not inkassa_page.context["entry_open"]

    filtered = client.get(reverse("cashbook:operations", kwargs={"desk": "bank"}) + "?q=DARVOZA&kind=income")
    assert filtered.context["active_filters"] == 2
    assert filtered.context["filters_open"]

    after_entry = client.get(reverse("cashbook:operations", kwargs={"desk": "bank"}) + "?entry_kind=income")
    assert after_entry.context["entry_open"]
