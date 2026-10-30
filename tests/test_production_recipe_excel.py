"""Recipes in Excel: the download, loading it back (also into an empty database) and the project's file."""

from decimal import Decimal
from io import BytesIO, StringIO

import openpyxl
import pytest
from django.contrib.messages import get_messages
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import CommandError, call_command
from django.urls import reverse
from django.utils import translation

from finance_saber.cashbook.exports import xlsx_bytes
from finance_saber.cashbook.handbooks import HandbookError
from finance_saber.production import recipe_excel
from finance_saber.production.models import Material, Product, ProductAlias, RecipeLine
from finance_saber.production.recipe_excel import load_recipes

PAGE = reverse("production:recipes")
URL = reverse("production:recipe_import")


@pytest.fixture
def recipe(db):
    """A product with a two-line recipe and another name."""
    product = Product.objects.create(name="B Fresh анор 0,450 л", sales_name="Анор", volume=Decimal("0.45"), position=3)
    sugar = Material.objects.create(name="Сахар", unit="кг", kind=Material.Kind.RAW, min_stock=100)
    can = Material.objects.create(name="Банка 0,450", unit="шт", kind=Material.Kind.PACKAGING)
    RecipeLine.objects.create(product=product, material=sugar, quantity_per_unit=Decimal("0.0475200000"), position=1)
    RecipeLine.objects.create(
        product=product, material=can, quantity_per_unit=1, allowed_loss=Decimal("0.5"), position=2
    )
    ProductAlias.objects.create(name="ЖБ 450 Мл Анор", product=product, source="SmartUp")
    return product


def workbook(language="ru") -> BytesIO:
    with translation.override(language):
        return BytesIO(xlsx_bytes(recipe_excel.recipe_sheets()))


def empty_database():
    RecipeLine.objects.all().delete()
    ProductAlias.objects.all().delete()
    Product.objects.all().delete()
    Material.objects.all().delete()


def edited(file: BytesIO, sheet: str, rows: list[list]) -> BytesIO:
    """``file`` with ``sheet``'s data rows replaced by ``rows``."""
    book = openpyxl.load_workbook(file)
    ws = book[sheet]
    ws.delete_rows(2, ws.max_row)
    for row in rows:
        ws.append(row)
    out = BytesIO()
    book.save(out)
    out.seek(0)
    return out


def test_download_has_the_four_sheets(client, user, recipe):
    client.force_login(user)
    page = client.get(PAGE).content.decode()
    assert reverse("production:recipe_export") in page and reverse("production:recipe_seed") in page
    response = client.get(reverse("production:recipe_export"))
    assert "attachment" in response["Content-Disposition"]
    book = openpyxl.load_workbook(BytesIO(response.content))
    assert book.sheetnames == ["Продукты", "Сырьё и упаковка", "Рецептуры", "Другие названия"]
    lines = list(book["Рецептуры"].iter_rows(min_row=2, values_only=True))
    assert lines[0][:5] == ("B Fresh анор 0,450 л", "Сахар", "кг", 0.04752, 105.6)  # per can and per 1000 l


@pytest.mark.parametrize("language", ["ru", "en"])
def test_the_download_loads_into_an_empty_database(recipe, language):
    file = workbook(language)
    empty_database()
    report = load_recipes(file)
    assert report.errors == [] and report.saved
    product = Product.objects.get()
    assert (product.sales_name, product.volume, product.position) == ("Анор", Decimal("0.450"), 3)
    assert Material.objects.get(name="Сахар").min_stock == 100
    assert Material.objects.get(name="Банка 0,450").kind == Material.Kind.PACKAGING
    norms = {line.material.name: (line.quantity_per_unit, line.allowed_loss) for line in product.recipe.all()}
    assert norms == {"Сахар": (Decimal("0.04752"), Decimal(2)), "Банка 0,450": (Decimal(1), Decimal("0.5"))}
    assert ProductAlias.objects.get().product == product

    again = load_recipes(workbook(language))
    assert not again.created and not again.updated and not again.removed


def test_dry_run_saves_nothing(recipe):
    file = workbook()
    empty_database()
    report = load_recipes(file, dry_run=True)
    assert report.errors == [] and not report.saved
    assert report.created["recipe_lines"] == 2
    assert not Product.objects.exists()


def test_the_file_sets_each_listed_recipe_exactly(recipe):
    other = Product.objects.create(name="Green Tea", volume=Decimal("0.45"))
    RecipeLine.objects.create(product=other, material=Material.objects.get(name="Сахар"), quantity_per_unit=1)
    # Only the anor recipe is in the file: sugar per 1000 l only, the can dropped.
    file = edited(workbook(), "Рецептуры", [["B Fresh анор 0,450 л", "Сахар", "кг", None, "110", 3, 1]])
    report = load_recipes(file)
    assert report.errors == [] and report.removed["recipe_lines"] == 1
    line = recipe.recipe.get()
    assert line.quantity_per_unit == Decimal("0.0495")  # 110 per 1000 l × 0.45 l
    assert line.allowed_loss == 3
    assert other.recipe.count() == 1  # not in the file: kept


def test_errors_save_nothing_and_name_the_row(recipe):
    file = edited(
        workbook(),
        "Рецептуры",
        [
            ["B Fresh анор 0,450 л", "Соль", "кг", 0.1],
            ["Нет такого", "Сахар", "кг", 0.1],
            ["B Fresh анор 0,450 л", "Сахар", "кг", "abc"],
        ],
    )
    with translation.override("ru"):
        report = load_recipes(file)
    assert not report.saved
    assert len(report.errors) == 3
    errors = " | ".join(report.errors)
    assert "строка 2: неизвестное сырьё «Соль»" in errors
    assert "строка 3: неизвестный продукт «Нет такого»" in errors
    assert "строка 4: «abc» — не число" in errors
    assert recipe.recipe.count() == 2  # unchanged


def test_not_a_recipes_file(db):
    with pytest.raises(HandbookError):
        load_recipes(BytesIO(b"not a workbook"))
    book = openpyxl.Workbook()
    out = BytesIO()
    book.save(out)
    out.seek(0)
    with pytest.raises(HandbookError, match="Рецептуры"):
        load_recipes(out)


def test_upload_page(client, technologist, recipe):
    file = workbook()
    empty_database()
    client.force_login(technologist)
    assert URL in client.get(PAGE).content.decode()
    upload = SimpleUploadedFile("recipes.xlsx", file.getvalue())
    response = client.post(URL, {"file": upload})
    assert response.context["report"].saved
    assert RecipeLine.objects.count() == 2


def test_viewers_cannot_load(client, finance, recipe):
    client.force_login(finance)
    assert URL not in client.get(PAGE).content.decode()
    assert client.get(reverse("production:recipe_export")).status_code == 200
    upload = SimpleUploadedFile("recipes.xlsx", workbook().getvalue())
    assert client.post(URL, {"file": upload}).status_code == 403


@pytest.fixture
def seed_file(tmp_path, monkeypatch, recipe):
    monkeypatch.setattr(recipe_excel, "SEED_DIR", tmp_path / "seed")
    return recipe_excel.write_seed_file("recipes")


def test_project_file_fills_an_empty_database_in_one_click(client, superuser, seed_file):
    empty_database()
    client.force_login(superuser)
    response = client.get(reverse("production:recipe_seed"))
    assert openpyxl.load_workbook(BytesIO(response.content)).sheetnames[2] == "Рецептуры"

    assert not client.post(URL, {"use_seed": "1", "dry_run": "on"}).context["report"].saved
    assert not Product.objects.exists()
    assert client.post(URL, {"use_seed": "1"}).context["report"].saved
    assert RecipeLine.objects.count() == 2


def test_project_file_is_not_public(client, no_role, seed_file):
    assert client.get(reverse("production:recipe_seed")).status_code == 302  # to sign-in
    client.force_login(no_role)
    assert client.get(reverse("production:recipe_seed")).status_code == 403


def test_missing_project_file(client, user, tmp_path, monkeypatch):
    monkeypatch.setattr(recipe_excel, "SEED_DIR", tmp_path / "none")
    client.force_login(user)
    response = client.get(reverse("production:recipe_seed"))
    assert response.status_code == 302
    assert any("export_recipes --seed" in str(m) for m in get_messages(response.wsgi_request))


def test_commands(recipe, tmp_path, monkeypatch):
    path = tmp_path / "recipes.xlsx"
    call_command("export_recipes", path, stdout=StringIO())
    empty_database()
    call_command("load_recipes", path, "--dry-run", stdout=StringIO())
    assert not Product.objects.exists()
    call_command("load_recipes", path, stdout=StringIO())
    assert RecipeLine.objects.count() == 2
    monkeypatch.setattr(recipe_excel, "SEED_DIR", tmp_path / "seed")
    call_command("export_recipes", "--seed", stdout=StringIO())
    assert (tmp_path / "seed" / "recipes.xlsx").is_file()
    with pytest.raises(CommandError, match="--seed"):
        call_command("export_recipes", stdout=StringIO())


def test_the_project_file_loads_into_the_test_database(db):
    """The file committed in production/seed/ is readable and loads without errors."""
    path = recipe_excel.seed_path("recipes")
    if not path.is_file():
        pytest.skip("the project has no recipes file")
    report = load_recipes(path, dry_run=True)
    assert report.errors == []
    assert report.created["recipe_lines"] > 0
