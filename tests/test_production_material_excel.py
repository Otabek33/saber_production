"""Materials in Excel: the download with stock, loading it back, opening balances and the project's file."""

import datetime
from decimal import Decimal
from io import BytesIO, StringIO

import openpyxl
import pytest
from django.contrib.messages import get_messages
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.urls import reverse
from django.utils import translation

from finance_saber.cashbook.exports import xlsx_bytes
from finance_saber.production import recipe_excel
from finance_saber.production.models import Material, MaterialMove, material_stock
from finance_saber.production.recipe_excel import MATERIALS, load_recipes

PAGE = reverse("production:materials")
URL = reverse("production:material_import")
OPENING = datetime.date(2026, 11, 1)


@pytest.fixture
def materials(db):
    """Sugar with 250 kg in stock, cans with none."""
    sugar = Material.objects.create(name="Сахар", unit="кг", kind=Material.Kind.RAW, min_stock=100, shelf_life="1 год")
    Material.objects.create(name="Банка 0,450", unit="шт", kind=Material.Kind.PACKAGING, is_active=False)
    MaterialMove.objects.create(
        material=sugar, date=datetime.date(2026, 10, 1), kind=MaterialMove.Kind.RECEIPT, quantity=250
    )
    return sugar


def workbook(language="ru") -> BytesIO:
    with translation.override(language):
        return BytesIO(xlsx_bytes(recipe_excel.material_sheets()))


def empty_database():
    MaterialMove.objects.all().delete()
    Material.objects.all().delete()


def load(file, **kwargs):
    return load_recipes(file, only=(MATERIALS,), **kwargs)


def test_download_has_the_stock(client, user, materials):
    client.force_login(user)
    page = client.get(PAGE).content.decode()
    for name in ("material_export", "material_seed", "material_import"):
        assert reverse(f"production:{name}") in page
    response = client.get(reverse("production:material_export"))
    book = openpyxl.load_workbook(BytesIO(response.content))
    assert book.sheetnames == ["Сырьё и упаковка"]
    rows = list(book.active.iter_rows(values_only=True))
    assert rows[0][-1] == "Остаток"
    assert rows[1][:3] == ("Сахар", "кг", "Сырьё") and rows[1][-1] == 250  # raw materials first


@pytest.mark.parametrize("language", ["ru", "en"])
def test_materials_load_into_an_empty_database_without_stock_by_default(materials, language):
    file = workbook(language)
    empty_database()
    report = load(file)
    assert report.errors == [] and report.saved and report.created["materials"] == 2
    sugar = Material.objects.get(name="Сахар")
    assert (sugar.unit, sugar.kind, sugar.min_stock, sugar.shelf_life) == ("кг", "raw", 100, "1 год")
    assert not Material.objects.get(name="Банка 0,450").is_active
    assert not MaterialMove.objects.exists()  # stock only when asked

    again = load(workbook(language))
    assert not again.created and not again.updated


def test_stock_as_opening_balances_on_an_empty_database(materials, superuser):
    file = workbook()
    empty_database()
    report = load(file, opening_date=OPENING, user=superuser)
    assert report.created["opening_balances"] == 1  # cans have no stock: nothing to open
    move = MaterialMove.objects.get()
    assert (move.kind, move.date, move.quantity, move.created_by) == ("opening", OPENING, 250, superuser)
    assert material_stock() == {move.material_id: 250}

    # Loading again changes nothing; corrected figures correct the opening balance, never add to it.
    assert load(workbook(), opening_date=OPENING).unchanged["opening_balances"] == 1
    book = openpyxl.load_workbook(workbook())
    book.active["G2"] = "240,5"
    corrected = BytesIO()
    book.save(corrected)
    corrected.seek(0)
    assert load(corrected, opening_date=OPENING).updated["opening_balances"] == 1
    assert material_stock() == {move.material_id: Decimal("240.5")}


def test_stock_is_never_loaded_over_existing_movements(materials):
    report = load(workbook(), opening_date=OPENING)
    assert report.stock_kept == ["Сахар"]
    assert material_stock() == {materials.pk: 250}  # unchanged, not doubled


def test_shortages_are_not_opening_balances(materials):
    file = workbook()
    empty_database()
    book = openpyxl.load_workbook(file)
    book.active["G2"] = -5
    out = BytesIO()
    book.save(out)
    out.seek(0)
    report = load(out, opening_date=OPENING)
    assert report.errors == [] and not MaterialMove.objects.exists()


def test_errors_save_nothing(materials):
    book = openpyxl.load_workbook(workbook())
    book.active.append(["Соль", "", "Непонятно"])
    out = BytesIO()
    book.save(out)
    out.seek(0)
    with translation.override("ru"):
        report = load(out)
    assert not report.saved and len(report.errors) == 1 and "строка 4" in report.errors[0]
    assert not Material.objects.filter(name="Соль").exists()


def test_the_materials_page_reads_only_materials(materials, superuser):
    """A recipes file uploaded on the materials page loads its materials, not its products."""
    from finance_saber.production.models import Product

    Product.objects.create(name="Чай", volume=Decimal("0.45"))
    with translation.override("ru"):
        file = BytesIO(xlsx_bytes(recipe_excel.recipe_sheets()))
    empty_database()
    Product.objects.all().delete()
    report = load(file)
    assert report.sheets == ["Сырьё и упаковка"] and not Product.objects.exists()


def test_upload_page_with_opening_balances(client, technologist, materials):
    file = workbook()
    empty_database()
    client.force_login(technologist)
    page = client.get(URL).content.decode()
    assert "opening_date" in page
    upload = SimpleUploadedFile("materials.xlsx", file.getvalue())
    response = client.post(URL, {"file": upload, "opening": "on", "opening_date": "2026-11-01"})
    assert response.context["report"].saved
    assert MaterialMove.objects.get().created_by == technologist


def test_upload_needs_a_file(client, technologist, materials):
    client.force_login(technologist)
    response = client.post(URL, {"dry_run": "on"})
    assert response.context["form"].errors["file"]


def test_viewers_cannot_load(client, finance, materials):
    client.force_login(finance)
    assert URL not in client.get(PAGE).content.decode()
    assert client.get(reverse("production:material_export")).status_code == 200
    upload = SimpleUploadedFile("materials.xlsx", workbook().getvalue())
    assert client.post(URL, {"file": upload}).status_code == 403


@pytest.fixture
def seed_dir(tmp_path, monkeypatch, materials):
    monkeypatch.setattr(recipe_excel, "SEED_DIR", tmp_path / "seed")
    recipe_excel.write_seed_file("materials")
    return tmp_path / "seed"


def test_project_file_fills_an_empty_database_in_one_click(client, superuser, seed_dir):
    empty_database()
    client.force_login(superuser)
    page = client.get(URL).content.decode()
    assert "materials.xlsx" in page and "use_seed" in page
    response = client.get(reverse("production:material_seed"))
    assert openpyxl.load_workbook(BytesIO(response.content)).sheetnames == ["Сырьё и упаковка"]

    assert not client.post(URL, {"use_seed": "1", "dry_run": "on"}).context["report"].saved
    assert not Material.objects.exists()
    response = client.post(URL, {"use_seed": "1", "opening": "on", "opening_date": "2026-11-01"})
    assert response.context["report"].saved
    assert Material.objects.count() == 2 and MaterialMove.objects.count() == 1


def test_project_file_is_not_public(client, no_role, seed_dir):
    assert client.get(reverse("production:material_seed")).status_code == 302  # to sign-in
    client.force_login(no_role)
    assert client.get(reverse("production:material_seed")).status_code == 403


def test_missing_project_file(client, user, tmp_path, monkeypatch):
    monkeypatch.setattr(recipe_excel, "SEED_DIR", tmp_path / "none")
    client.force_login(user)
    response = client.get(reverse("production:material_seed"))
    assert response.status_code == 302
    assert any("export_materials --seed" in str(m) for m in get_messages(response.wsgi_request))


def test_commands(materials, tmp_path, monkeypatch):
    path = tmp_path / "materials.xlsx"
    call_command("export_materials", path, stdout=StringIO())
    empty_database()
    call_command("load_materials", path, stdout=StringIO())
    assert Material.objects.count() == 2 and not MaterialMove.objects.exists()
    call_command("load_materials", path, "--opening-date", "2026-11-01", stdout=StringIO())
    assert MaterialMove.objects.get().date == OPENING
    monkeypatch.setattr(recipe_excel, "SEED_DIR", tmp_path / "seed")
    call_command("export_materials", "--seed", stdout=StringIO())
    assert (tmp_path / "seed" / "materials.xlsx").is_file()


def test_the_project_file_loads_into_the_test_database(db):
    """The file committed in production/seed/ is readable and loads, with its stock, without errors."""
    path = recipe_excel.seed_path("materials")
    if not path.is_file():
        pytest.skip("the project has no materials file")
    report = load(path, dry_run=True, opening_date=OPENING)
    assert report.errors == [] and report.created["materials"] > 0
