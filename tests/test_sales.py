"""Sales: daily entry, duplicates, stock, debts, SmartUp import and the yearly plan."""

import datetime
from decimal import Decimal

import openpyxl
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse

from finance_saber.cashbook.models import Account, CashDesk, Category, Counterparty, Kind, Operation
from finance_saber.production.models import Material, MaterialMove, Product, ProductMove, product_stock
from finance_saber.sales import services
from finance_saber.sales.models import Channel, Client, Plan, Sale

D = Decimal
DAY = datetime.date(2026, 10, 23)


@pytest.fixture
def tea(db):
    return Product.objects.create(
        name="Green Tea мёд 0,450 л", sales_name="Sakura Зеленый чай с мёдом", volume=D("0.45")
    )


@pytest.fixture
def energy(db):
    return Product.objects.create(name="7 Energy Drink 0,450 л", sales_name="7 Energy 450 ML 1/12", volume=D("0.45"))


@pytest.fixture
def client_shop(db):
    return Client.objects.create(name="M094-ТАШ..ЭКС.ЮНУСОБОД 6КВ")


@pytest.fixture
def stock_counted(db):
    """Finished goods are tracked from 22.10 (the warehouse's opening count)."""
    material = Material.objects.create(name="Сахар", unit="кг")
    MaterialMove.objects.create(
        material=material, date=datetime.date(2026, 10, 22), kind=MaterialMove.Kind.OPENING, quantity=1
    )


def sale(client, product, amount, date=DAY, quantity=12):
    return services.record_sale(Sale(date=date, client=client, product=product, quantity=quantity, amount=D(amount)))


# --- Recording ---------------------------------------------------------------------------


def test_sale_takes_goods_out_of_stock(stock_counted, client_shop, tea):
    sale(client_shop, tea, "110004")
    assert product_stock() == {tea.pk: D("-12")}
    move = ProductMove.objects.get()
    assert (move.kind, move.source) == (ProductMove.Kind.SHIPMENT, f"sale:{Sale.objects.get().pk}")


def test_sales_before_the_stock_count_are_history(stock_counted, client_shop, tea):
    old = sale(client_shop, tea, "110004", date=datetime.date(2026, 10, 3))
    assert not old.affects_stock
    assert not ProductMove.objects.exists()


def test_cancelling_a_sale_returns_the_goods(stock_counted, client_shop, tea):
    recorded = sale(client_shop, tea, "110004")
    services.set_cancelled(recorded, True)
    assert product_stock().get(tea.pk, 0) == 0
    services.set_cancelled(recorded, False)
    assert product_stock()[tea.pk] == D("-12")


def test_restoring_is_refused_when_it_would_duplicate(client_shop, tea):
    first = sale(client_shop, tea, "110004")
    services.set_cancelled(first, True)
    sale(client_shop, tea, "110004")
    with pytest.raises(ValueError):
        services.set_cancelled(first, False)


def test_same_day_client_product_and_amount_is_a_duplicate(client_shop, tea, energy):
    sale(client_shop, tea, "110004")
    assert services.is_duplicate(DAY, client_shop, tea, D("110004"))
    # A different product for the same amount is a different sale (SmartUp has many of these).
    assert not services.is_duplicate(DAY, client_shop, energy, D("110004"))


def test_the_database_refuses_duplicates_too(client_shop, tea):
    from django.db import IntegrityError, transaction

    sale(client_shop, tea, "110004")
    with pytest.raises(IntegrityError), transaction.atomic():
        Sale.objects.create(date=DAY, client=client_shop, product=tea, quantity=1, amount=D("110004"))


# --- Daily page --------------------------------------------------------------------------


def post_sale(http, **data):
    form = {"date": DAY.isoformat(), "quantity": "12", "amount": "110 004", "note": "", "channel": ""} | data
    return http.post(reverse("sales:daily") + f"?date={DAY.isoformat()}", form)


def test_daily_page_adds_a_sale_for_an_existing_client(client, user, client_shop, tea):
    client.force_login(user)
    response = post_sale(client, client=str(client_shop.pk), product=tea.pk)
    assert response.status_code == 302
    recorded = Sale.objects.get()
    assert (recorded.client, recorded.amount, recorded.created_by) == (client_shop, D("110004"), user)


def test_typing_a_new_client_name_creates_the_client(client, user, tea):
    client.force_login(user)
    post_sale(client, client="Новый магазин", product=tea.pk)
    assert Sale.objects.get().client.name == "Новый магазин"


def test_invalid_entry_creates_no_client(client, user, tea):
    client.force_login(user)
    response = post_sale(client, client="Новый магазин", product=tea.pk, quantity="")
    assert response.status_code == 200
    assert not Client.objects.exists()


def test_daily_page_refuses_a_duplicate(client, user, client_shop, tea):
    sale(client_shop, tea, "110004")
    client.force_login(user)
    response = post_sale(client, client=str(client_shop.pk), product=tea.pk)
    assert response.status_code == 200
    assert response.context["form"].has_error("__all__", "duplicate")
    assert Sale.objects.count() == 1


def test_daily_page_lists_the_day(client, user, client_shop, tea, energy):
    sale(client_shop, tea, "110004")
    sale(client_shop, energy, "220008", quantity=24)
    client.force_login(user)
    response = client.get(reverse("sales:daily"), {"date": DAY.isoformat()})
    assert response.context["total_quantity"] == 36
    assert response.context["total_amount"] == D("330012")


def test_cancel_and_edit(client, user, client_shop, tea):
    recorded = sale(client_shop, tea, "110004")
    client.force_login(user)
    client.post(reverse("sales:sale_cancel", args=[recorded.pk]))
    recorded.refresh_from_db()
    assert recorded.is_cancelled
    client.post(reverse("sales:sale_cancel", args=[recorded.pk]))
    response = client.post(
        reverse("sales:sale_edit", args=[recorded.pk]),
        {
            "date": DAY.isoformat(),
            "client": str(client_shop.pk),
            "product": tea.pk,
            "quantity": "24",
            "amount": "220008",
        },
    )
    assert response.status_code == 302
    recorded.refresh_from_db()
    assert (recorded.quantity, recorded.amount, recorded.is_cancelled) == (24, D("220008"), False)


# --- Debts -------------------------------------------------------------------------------


@pytest.fixture
def payer(db):
    return Counterparty.objects.create(reference="2020/304577774/CHAPAN", inn="304577774", name="CHAPAN")


def bank_payment(payer, amount, date=DAY, kind=Kind.INCOME):
    desk = CashDesk.objects.get(slug="bank")
    category = Category.objects.first()
    return Operation.objects.create(
        desk=desk, kind=kind, date=date, category=category, account=Account.BANK, amount=D(amount), counterparty=payer
    )


def test_client_without_payer_owes_its_shipments(client_shop, tea):
    sale(client_shop, tea, "110004")
    [row] = services.debt_rows()
    assert (row.shipped, row.paid, row.balance) == (D("110004"), 0, D("110004"))


def test_payer_payments_settle_all_its_clients(client_shop, tea, payer):
    other = Client.objects.create(name="M134-ТАШ..ВОДНИК CHAPAN", payer=payer)
    client_shop.payer = payer
    client_shop.save()
    sale(client_shop, tea, "110004")
    sale(other, tea, "110004")
    bank_payment(payer, "300000")
    bank_payment(payer, "5000", date=DAY - datetime.timedelta(days=30))  # before the first sale: an older debt
    [row] = services.debt_rows()
    assert (row.name, len(row.clients), row.shipped, row.paid) == ("CHAPAN", 2, D("220008"), D("300000"))
    assert row.balance == D("-79992")  # paid in advance: a creditor


def test_refunds_and_opening_balance_count(client_shop, tea, payer):
    client_shop.payer, client_shop.opening_balance = payer, D("1000")
    client_shop.save()
    sale(client_shop, tea, "110004")
    bank_payment(payer, "110004")
    bank_payment(payer, "500", kind=Kind.EXPENSE)
    [row] = services.debt_rows()
    assert row.balance == D("1500")


def test_statement_runs_the_balance(client_shop, tea, payer):
    client_shop.payer = payer
    client_shop.save()
    sale(client_shop, tea, "110004")
    bank_payment(payer, "100000", date=DAY + datetime.timedelta(days=1))
    lines = services.statement([client_shop], payer=payer)
    assert [(line.kind, line.balance) for line in lines] == [("sale", D("110004")), ("payment", D("10004"))]


def test_debts_and_statement_pages(client, user, client_shop, tea, payer):
    sale(client_shop, tea, "110004")
    client.force_login(user)
    response = client.get(reverse("sales:debts"))
    assert response.context["receivable"] == D("110004")
    assert client.get(reverse("sales:statement", args=["client", client_shop.pk])).status_code == 200
    client_shop.payer = payer
    client_shop.save()
    assert client.get(reverse("sales:statement", args=["client", client_shop.pk])).status_code == 302


def test_linking_a_payer_on_the_client_page(client, finance, client_shop, payer):
    client.force_login(finance)
    response = client.post(
        reverse("sales:client_edit", args=[client_shop.pk]),
        {"name": client_shop.name, "payer": payer.pk, "opening_balance": "-1 000", "is_active": "on"},
    )
    assert response.status_code == 302
    client_shop.refresh_from_db()
    assert (client_shop.payer, client_shop.opening_balance) == (payer, D("-1000"))


# --- SmartUp export ----------------------------------------------------------------------


def smartup_file(tmp_path, rows):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "SmartUp"
    sheet.append(
        [None, "Дата отгрузки", "Организация", "Название проекта", "ТМЦ", "Клиент", "Количество продажи",
         "Сумма с переоценкой в базовой валюте", None, None, None, None, None, "Название проекта"]
    )  # fmt: skip
    for row in rows:
        sheet.append(row)
    sheet["N2"], sheet["O2"] = "Export", "экс"
    path = tmp_path / "smartup.xlsx"
    book.save(path)
    return path


SMARTUP_ROWS = [
    [
        None,
        "23.10.2026",
        "BON NATURAL DRINK",
        "Key Account",
        "Sakura Зеленый чай с мёдом",
        "M094 магазин",
        12,
        110004.0,
    ],
    [None, "23.10.2026", "BON NATURAL DRINK", "Key Account", "7 Energy 450 ML 1/12", "M094 магазин", 24, 110004.0],
    [None, "23.10.2026", "BON NATURAL DRINK", "Export", "7 Energy 450 ML 1/12", "Kazakh LLP", 2400, 9000000],
    [None, "23.10.2026", "BON NATURAL DRINK", "Key Account", "7 Energy 450 ML 1/12", "M083 магазин", None, None],
]


def test_read_smartup(tmp_path):
    parsed = services.read_smartup(smartup_file(tmp_path, SMARTUP_ROWS))
    assert len(parsed.rows) == 3  # the empty line is skipped
    assert parsed.rows[0].date == DAY
    assert parsed.export_channels == {"Export"}


def test_import_sales(tmp_path, user, tea, energy):
    report = services.import_sales(services.read_smartup(smartup_file(tmp_path, SMARTUP_ROWS)), user=user)
    assert report.saved and report.created == 3
    assert sorted(report.new_clients) == ["Kazakh LLP", "M094 магазин"]
    assert Channel.objects.get(name="Export").is_export
    # The same file again: everything is a duplicate.
    again = services.import_sales(services.read_smartup(smartup_file(tmp_path, SMARTUP_ROWS)), user=user)
    assert (again.created, len(again.duplicates)) == (0, 3)


def test_unknown_product_stops_the_import(tmp_path, tea):
    rows = [*SMARTUP_ROWS[:1], [None, "23.10.2026", "X", "Key Account", "Неизвестный напиток", "M094", 12, 1000]]
    report = services.import_sales(services.read_smartup(smartup_file(tmp_path, rows)))
    assert report.unknown_products == ["Неизвестный напиток"]
    assert not report.saved and not Sale.objects.exists()


def test_known_aliases_match_smartup_names(tmp_path, energy):
    rows = [[None, "23.10.2026", "X", "Key Account", "7 Energy 450 ML  1/24", "M094", 24, 220008]]
    report = services.import_sales(services.read_smartup(smartup_file(tmp_path, rows)))
    assert report.saved
    assert Sale.objects.get().product == energy


def test_import_page(client, user, tmp_path, tea, energy):
    client.force_login(user)
    upload = SimpleUploadedFile("day.xlsx", smartup_file(tmp_path, SMARTUP_ROWS).read_bytes())
    response = client.post(reverse("sales:import"), {"sales-file": upload})
    assert response.context["sales_report"].created == 3
    assert Sale.objects.count() == 3


# --- Plan --------------------------------------------------------------------------------


def plan_file(tmp_path):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Plan 2026"
    sheet["A1"] = 2026
    for title_row, title, values in (
        (4, "Всего план реализации на 2026 год", 100),
        (27, "Всего план производства на 2026 год", 150),
    ):
        sheet.cell(title_row, 2, title)
        sheet.cell(title_row + 1, 3, "Наименование товара")
        for month in range(1, 13):
            sheet.cell(title_row + 1, 3 + month, datetime.datetime(2026, month, 1))
            sheet.cell(title_row + 2, 3 + month, "План")
        for offset, name in enumerate(["Sakura Зеленый чай с мёдом", "7 Energy 450 ML 1/12"]):
            sheet.cell(title_row + 3 + offset, 3, name)
            for month in range(1, 13):
                sheet.cell(title_row + 3 + offset, 3 + month, values * month)
    path = tmp_path / "plan.xlsx"
    book.save(path)
    return path


def test_import_plan(tmp_path, tea, energy):
    report = services.import_plan(services.read_plan(plan_file(tmp_path)))
    assert report.saved and report.year == 2026
    assert Plan.objects.filter(kind=Plan.Kind.SALES).count() == 24
    assert Plan.objects.get(kind=Plan.Kind.PRODUCTION, product=tea, month=3).quantity == 450
    services.import_plan(services.read_plan(plan_file(tmp_path)))  # again: replaces, doesn't add
    assert Plan.objects.count() == 48


def test_plan_page_shows_plan_and_fact(client, user, tmp_path, client_shop, tea, energy):
    services.import_plan(services.read_plan(plan_file(tmp_path)))
    sale(client_shop, tea, "110004", date=datetime.date(2026, 10, 5), quantity=500)
    client.force_login(user)
    response = client.get(reverse("sales:plan"), {"year": 2026})
    row = next(row for row in response.context["rows"] if row["product"] == tea)
    october = row["cells"][9]
    assert (october["plan"], october["fact"], october["percent"]) == (1000, 500, 50)


def test_plan_can_be_edited(client, planner, tea):
    client.force_login(planner)
    response = client.post(reverse("sales:plan") + "?year=2026&kind=sales", {f"plan-{tea.pk}-11": "1 500"})
    assert response.status_code == 302
    assert Plan.objects.get(product=tea, month=11, year=2026, kind=Plan.Kind.SALES).quantity == 1500


# --- Roles ------------------------------------------------------------------------------


def test_planner_reads_sales_but_cannot_enter_them(client, planner, tea, client_shop):
    client.force_login(planner)
    assert client.get(reverse("sales:daily")).status_code == 200
    assert client.get(reverse("sales:debts")).status_code == 200
    assert post_sale(client, client=str(client_shop.pk), product=tea.pk).status_code == 403
    assert client.get(reverse("sales:import")).status_code == 403


def test_no_role_sees_no_sales(client, no_role):
    client.force_login(no_role)
    for name in ("sales:daily", "sales:debts", "sales:plan", "sales:clients"):
        assert client.get(reverse(name)).status_code == 403
