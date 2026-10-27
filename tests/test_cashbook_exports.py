"""Excel downloads of the handbook and the counterparties list."""

from decimal import Decimal
from io import BytesIO
from urllib.parse import quote

import openpyxl
from django.urls import reverse
from test_cashbook_bank import NEW, ROWS, SALES, counterparties, desk  # noqa: F401 (fixtures)

from finance_saber.cashbook import bank
from finance_saber.cashbook.models import Article, Counterparty

HANDBOOK = reverse("cashbook:handbook_export")
COUNTERPARTIES = reverse("cashbook:counterparty_export")


def workbook_of(response):
    assert response["Content-Type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return openpyxl.load_workbook(BytesIO(response.content))


def rows(sheet):
    return [list(values) for values in sheet.iter_rows(values_only=True)]


def test_handbook_download(client, user):
    client.force_login(user)
    assert HANDBOOK in client.get(reverse("cashbook:handbook")).content.decode()

    response = client.get(HANDBOOK)
    assert response.status_code == 200
    assert "attachment" in response["Content-Disposition"]
    assert quote("Справочник") in response["Content-Disposition"]
    book = workbook_of(response)
    assert book.sheetnames == ["Статьи и категории", "Статьи"]

    flat = rows(book["Статьи и категории"])
    assert flat[0] == [
        "Тип", "Статья", "Категория", "Кассы", "Учитывается в ДДС", "Статья активна", "Категория активна",
    ]  # fmt: skip
    salary = next(row for row in flat if row[1] == "Взаиморасчеты с персоналом" and row[2] == "ЗП")
    assert salary[4:] == ["да", "да", "да"]
    assert book["Статьи и категории"].freeze_panes == "A2"

    articles = rows(book["Статьи"])
    assert len(articles) == Article.objects.count() + 1
    assert articles[0][-1] == "Категорий"


def test_counterparties_download_keeps_the_page_filters(client, user, desk, counterparties):  # noqa: F811
    bank.record_lines(desk, bank.read_statement(BytesIO(statement_bytes())))
    client.force_login(user)
    page = client.get(reverse("cashbook:counterparties") + "?status=unclassified").content.decode()
    assert f"{COUNTERPARTIES}?status=unclassified" in page

    everyone = rows(workbook_of(client.get(COUNTERPARTIES)).active)
    assert everyone[0][:4] == ["Название", "ИНН", "Номер счёта", "Статья"]
    assert len(everyone) == 1 + Counterparty.objects.count()
    sales = next(row for row in everyone if row[-1] == SALES)
    sale_kind = str(Article.objects.get(name="Месный продажа").get_kind_display())
    assert sales[1:8] == ["304577774", "20208000900713230001", "Месный продажа", sale_kind, "Tash reg", None, 2]
    assert (Decimal(str(sales[8])), Decimal(str(sales[9]))) == (Decimal("72000000"), Decimal("1500000"))

    unclassified = rows(workbook_of(client.get(COUNTERPARTIES + "?status=unclassified")).active)
    assert [row[-1] for row in unclassified[1:]] == [NEW]
    assert unclassified[1][3] == "Нужна статья"
    assert client.get(COUNTERPARTIES + "?region=abc").status_code == 200  # a bad filter is ignored, not an error


def test_text_that_looks_like_a_formula_stays_text(client, user):
    Counterparty.objects.create(reference="=HYPERLINK(1)", name="=HYPERLINK(1)")
    client.force_login(user)
    sheet = workbook_of(client.get(COUNTERPARTIES)).active
    cell = sheet["A2"]
    assert (cell.value, cell.data_type) == ("=HYPERLINK(1)", "s")


def test_downloads_need_the_page_permissions(client, no_role):
    client.force_login(no_role)
    assert client.get(HANDBOOK).status_code == 403
    assert client.get(COUNTERPARTIES).status_code == 403


def statement_bytes():
    from test_cashbook_bank import statement

    return statement(ROWS)
