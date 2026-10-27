"""The "Upload Excel" page: the finance workbook checked, then imported into every cash desk."""

import datetime
from io import BytesIO

import openpyxl
import pytest
from django.contrib.messages import get_messages
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from test_cashbook_bank import NEW, ROWS, counterparties  # noqa: F401 (fixture)
from test_cashbook_import import DAYS_ROWS, HEADERS, TAXI

from finance_saber.cashbook import workbook
from finance_saber.cashbook.models import CashDesk, Counterparty, Operation

URL = reverse("cashbook:workbook_import")
DAY = datetime.datetime(2026, 10, 20)
BANK_HEADERS = ["Дата", "Дата", "Cчет/ИНН", "№ док", "Оп", "МФО", "Оборот Дебет", "Оборот Кредит", "Назначение платежа"]
INKASSA_ROWS = [{"B": "Приход розница", "D": DAY, "E": "розница", "F": 500000, "H": 1500000}]


def finance_workbook(bon_rows=DAYS_ROWS, inkassa_rows=INKASSA_ROWS, bank_rows=ROWS, sheets=None):
    """Situation.xlsx with its three cash sheets, as an uploaded file."""
    book = openpyxl.Workbook()
    book.remove(book.active)
    if bon_rows is not None:
        bon = book.create_sheet("Касса Бон")
        for letter, text in HEADERS.items():
            bon[f"{letter}4"] = text
        for number, values in enumerate(bon_rows, start=6):
            for letter, value in values.items():
                bon[f"{letter}{number}"] = value
    if inkassa_rows is not None:
        inkassa = book.create_sheet("Инкасса & Розница")
        for letter, text in {"B": "Категория", "D": "Дата", "E": "Примичание", "F": "UZS", "G": "UZS"}.items():
            inkassa[f"{letter}4"] = text
        inkassa["H5"] = 1000000  # opening balance
        for number, values in enumerate(inkassa_rows, start=6):
            for letter, value in values.items():
                inkassa[f"{letter}{number}"] = value
    if bank_rows is not None:
        bank_sheet = book.create_sheet("Банк")
        bank_sheet.append(["Выписка"])
        bank_sheet.append(BANK_HEADERS)
        for posted_at, reference, document, debit, credit, purpose in bank_rows:
            bank_sheet.append(
                [posted_at.date(), posted_at, reference, document, 4 if credit else 1, "00440", debit, credit, purpose]
            )
    for name in sheets or ():
        book.create_sheet(name)
    buffer = BytesIO()
    book.save(buffer)
    return SimpleUploadedFile("Situation (1).xlsx", buffer.getvalue())


@pytest.fixture(autouse=True)
def upload_dir(tmp_path, monkeypatch):
    monkeypatch.setattr(workbook, "UPLOAD_DIR", tmp_path / "uploads")


def results(response):
    return {row.desk.slug: row for row in response.context["report"].results}


def messages_of(response):
    return [str(message) for message in get_messages(response.wsgi_request)]


def test_check_then_import(client, user, counterparties):  # noqa: F811
    client.force_login(user)
    page = client.get(URL).content.decode()
    assert "Финансовая книга (.xlsx)" in page

    response = client.post(URL, {"check": "", "file": finance_workbook()})
    checked = results(response)
    assert (checked["bon"].created, checked["inkassa"].created, checked["bank"].created) == (5, 2, 5)
    assert checked["bank"].new_counterparties == ["TOSHKENT XALQARO AEROPORTI MCHJ"]
    assert not Operation.objects.exists()  # only checked
    assert not Counterparty.objects.filter(reference=NEW).exists()
    content = response.content.decode()
    assert "1 контрагента ещё нет в программе" in content
    assert "Проверено, ещё не сохранено" in content

    response = client.post(URL, {"import": "", "counterparties": "add"})
    imported = results(response)
    assert response.context["result"].saved
    assert imported["bon"].url.endswith("/cash/bon/?date_from=2026-10-20&date_to=2026-10-20")
    assert Operation.objects.filter(desk__slug="bon").count() == 5
    assert Operation.objects.filter(desk__slug="inkassa").count() == 2
    assert Operation.objects.filter(desk__slug="bank").count() == 5
    assert Counterparty.objects.filter(reference=NEW).exists()
    assert "Книга загружена." in messages_of(response)
    assert not list(workbook.UPLOAD_DIR.glob("*.xlsx"))  # the file is deleted after the import

    # The same workbook again: nothing new, nothing duplicated.
    response = client.post(URL, {"check": "", "file": finance_workbook()})
    assert all(row.created == 0 for row in results(response).values())
    assert "Всё из этой книги уже загружено" in response.content.decode()


def test_new_counterparties_can_be_left_for_later(client, user, counterparties):  # noqa: F811
    client.force_login(user)
    client.post(URL, {"check": "", "file": finance_workbook()})
    response = client.post(URL, {"import": "", "counterparties": "skip"})
    bank = results(response)["bank"]
    assert (bank.created, bank.skipped_unknown) == (3, 2)
    assert not Counterparty.objects.filter(reference=NEW).exists()
    assert "строк новых контрагентов пропущено: 2" in response.content.decode()

    # Once the counterparty is added by hand, the next upload brings its lines.
    Counterparty.objects.create(reference=NEW, **Counterparty.parse_reference(NEW))
    client.post(URL, {"check": "", "file": finance_workbook()})
    response = client.post(URL, {"import": "", "counterparties": "skip"})
    assert results(response)["bank"].created == 2
    assert Operation.objects.filter(desk__slug="bank").count() == 5


def test_a_section_with_errors_is_skipped_and_the_others_imported(client, user, counterparties):  # noqa: F811
    client.force_login(user)
    client.post(URL, {"check": "", "file": finance_workbook(bank_rows=ROWS[:1])})
    client.post(URL, {"import": ""})

    changed = [dict(values) for values in DAYS_ROWS]
    changed[1]["N"] = 1900000  # the business trip was 1 850 000 when imported
    response = client.post(URL, {"check": "", "file": finance_workbook(bon_rows=[*changed, TAXI])})
    checked = results(response)
    assert not checked["bon"].ok
    assert [conflict.row for conflict in checked["bon"].conflicts] == [7]
    assert "строка 7: здесь 20.10.2026" in response.content.decode()

    response = client.post(URL, {"import": ""})
    assert not results(response)["bon"].ok
    assert Operation.objects.filter(desk__slug="bon").count() == 5  # not even the new taxi row
    assert Operation.objects.filter(desk__slug="bank").count() == 5  # the bank's other 4 lines came in
    assert "Загружено, кроме разделов с ошибками (см. ниже)." in messages_of(response)


def test_a_missing_sheet_is_named(client, user, counterparties):  # noqa: F811
    client.force_login(user)
    response = client.post(URL, {"check": "", "file": finance_workbook(bank_rows=None, sheets=["Лист1"])})
    bank = results(response)["bank"]
    assert bank.error == "Лист «Банк» не найден. В книге есть: «Касса Бон», «Инкасса & Розница», «Лист1»."
    assert results(response)["bon"].ok
    response = client.post(URL, {"import": ""})
    assert not Operation.objects.filter(desk__slug="bank").exists()
    assert Operation.objects.filter(desk__slug="bon").count() == 5


def test_import_without_a_checked_file(client, user):
    client.force_login(user)
    response = client.post(URL, {"import": ""})
    assert response.status_code == 302
    assert "Проверенный файл больше недоступен. Загрузите и проверьте его снова." in messages_of(response)


def test_not_a_workbook(client, user):
    client.force_login(user)
    response = client.post(URL, {"check": "", "file": SimpleUploadedFile("Situation.xlsx", b"not a workbook")})
    assert all(not row.ok for row in results(response).values())
    assert not Operation.objects.exists()


def test_menu_and_desk_pages_link_to_the_page(client, user):
    client.force_login(user)
    for desk in CashDesk.objects.filter(is_active=True):
        assert URL in client.get(desk.get_absolute_url()).content.decode()


def test_only_people_who_enter_operations_can_upload(client, no_role):
    client.force_login(no_role)
    assert client.get(URL).status_code == 403


def test_nothing_new_names_the_last_dates_and_keeps_no_file(client, user, counterparties):  # noqa: F811
    """A workbook with no later rows (e.g. not saved in Excel) is explained, not just "not saved"."""
    client.force_login(user)
    client.post(URL, {"check": "", "file": finance_workbook()})
    client.post(URL, {"import": ""})

    response = client.post(URL, {"check": "", "file": finance_workbook()})
    checked = results(response)
    assert checked["bon"].last_date_in_file == checked["bon"].last_date_here == DAY.date()
    content = response.content.decode()
    assert "Касса БОН — 20.10.2026" in content
    assert "сохраните книгу в Excel (Ctrl+S)" in content
    assert 'name="import"' not in content
    assert not list(workbook.UPLOAD_DIR.glob("*.xlsx"))  # nothing to import, so the file isn't kept
    assert "cashbook_workbook" not in client.session  # and no stale token in the session


def test_changed_rows_can_be_updated_from_excel(client, user, counterparties):  # noqa: F811
    """Excel is the official record during the trial: a corrected amount there can be taken over."""
    client.force_login(user)
    client.post(URL, {"check": "", "file": finance_workbook()})
    client.post(URL, {"import": ""})

    changed = [dict(values) for values in DAYS_ROWS]
    changed[1]["N"] = 1900000  # the business trip was 1 850 000 when imported
    response = client.post(URL, {"check": "", "file": finance_workbook(bon_rows=[*changed, TAXI])})
    bon = results(response)["bon"]
    # Shown as waiting though the section is blocked: the next day's taxi, and a correction for row 10's
    # hand-typed card balance, which no longer matches the corrected trip.
    assert bon.created == 2
    content = response.content.decode()
    assert "Изменены строки" in content
    assert "Обновить эти строки по Excel" in content

    response = client.post(URL, {"import": "", "update_changed": "on"})
    bon = results(response)["bon"]
    assert bon.ok
    assert [change.row for change in bon.updated] == [7]
    assert bon.created == 2
    trip = Operation.objects.get(source="Situation.xlsx/Касса Бон/7/1")
    assert trip.amount == 1900000
    assert Operation.objects.filter(desk__slug="bon").count() == 7
    assert "строка 7: было 20.10.2026" in response.content.decode()

    # Now Excel and the app agree: the next check finds nothing.
    response = client.post(URL, {"check": "", "file": finance_workbook(bon_rows=[*changed, TAXI])})
    assert not results(response)["bon"].conflicts
