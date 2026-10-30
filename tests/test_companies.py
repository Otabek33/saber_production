from io import BytesIO, StringIO

import openpyxl
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.db import IntegrityError
from django.urls import reverse
from django.utils.html import escape

from finance_saber.cashbook import company_files, exports
from finance_saber.cashbook.handbooks import load_files
from finance_saber.cashbook.models import Company, Counterparty, Region


@pytest.fixture
def own(db):
    return Company.objects.create(name='ООО "BON NATURAL DRINK"', inn="305900839")


@pytest.fixture
def accounts(own):
    connected = Counterparty.objects.create(
        reference="20208000700950671001/305900839/OOO BON NATURAL DRINK",
        name='OOO "BON NATURAL DRINK"',
        inn="305900839",
        account_number="20208000700950671001",
        company=own,
    )
    loose = Counterparty.objects.create(
        reference="40702000000000000001/770000001/ООО Ромашка", name="ООО Ромашка", inn="770000001"
    )
    return connected, loose


def test_page_lists_each_company_with_its_accounts(client, finance, own, accounts):
    client.force_login(finance)
    response = client.get(reverse("cashbook:companies"))
    assert response.status_code == 200
    page = response.content.decode()
    assert escape(own.name) in page
    assert "20208000700950671001" in page
    assert "ООО Ромашка" not in page  # not connected: counted, not listed under a company
    assert response.context["unconnected_count"] == 1


def test_company_is_created_with_a_unique_inn(client, finance, own):
    client.force_login(finance)
    response = client.post(reverse("cashbook:company_add"), {"name": "ООО Вторая", "inn": "123456789"})
    assert response.status_code == 302
    assert Company.objects.filter(inn="123456789").exists()

    response = client.post(reverse("cashbook:company_add"), {"name": "Дубль", "inn": own.inn})
    assert response.status_code == 200
    assert Company.objects.filter(inn=own.inn).count() == 1


def test_bank_account_is_connected_from_its_edit_form(client, finance, own, accounts):
    client.force_login(finance)
    _, loose = accounts
    response = client.post(
        reverse("cashbook:counterparty_edit", args=[loose.pk]),
        {"name": loose.name, "inn": loose.inn, "account_number": "", "company": own.pk},
    )
    assert response.status_code == 302
    loose.refresh_from_db()
    assert loose.company == own


def test_counterparties_filter_by_company(client, finance, own, accounts):
    client.force_login(finance)
    url = reverse("cashbook:counterparties")
    connected_only = client.get(url, {"company": own.pk}).content.decode()
    assert "20208000700950671001" in connected_only
    assert "ООО Ромашка" not in connected_only
    not_connected = client.get(url, {"company": "none"}).content.decode()
    assert "ООО Ромашка" in not_connected
    assert "20208000700950671001" not in not_connected


def test_deleting_a_company_keeps_its_accounts(own, accounts):
    connected, _ = accounts
    own.delete()
    connected.refresh_from_db()
    assert connected.company is None


def test_inn_is_unique_at_database_level(own):
    with pytest.raises(IntegrityError):
        Company.objects.create(name="Other", inn=own.inn)


def workbook_bytes(rows):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["Company", "INN", "Our company", "Statement reference"])
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def upload(client, rows, dry_run=False):
    data = {"file": SimpleUploadedFile("companies.xlsx", workbook_bytes(rows)), "dry_run": dry_run}
    return client.post(reverse("cashbook:companies_import"), data)


def test_download_has_one_row_per_account(client, finance, own, accounts):
    client.force_login(finance)
    response = client.get(reverse("cashbook:companies_export"))
    assert response.status_code == 200
    assert response["Content-Type"].startswith("application/vnd.openxmlformats")
    sheet = openpyxl.load_workbook(BytesIO(response.content)).active
    rows = [list(row) for row in sheet.iter_rows(min_row=2, values_only=True)]
    assert len(rows) == 1
    name, inn, region, is_own, reference = rows[0]
    assert (name, inn, region, reference) == (own.name, own.inn, None, accounts[0].reference)
    assert is_own in ("no", "нет")


def test_upload_dry_run_saves_nothing(client, finance):
    client.force_login(finance)
    upload(client, [["ООО Новая", "123456789", "", ""]], dry_run=True)
    assert not Company.objects.filter(inn="123456789").exists()


def test_upload_creates_company_and_connects_its_account(client, finance, accounts):
    client.force_login(finance)
    connected, loose = accounts
    upload(client, [["ООО Ромашка", loose.inn, "", loose.reference]])
    loose.refresh_from_db()
    assert loose.company is not None
    assert loose.company.inn == loose.inn
    assert loose.company.name == "ООО Ромашка"


def test_upload_with_a_bad_reference_saves_nothing(client, finance, accounts):
    client.force_login(finance)
    _, loose = accounts
    upload(client, [["ООО Ромашка", loose.inn, "", loose.reference], ["Нет", "999000111", "", "no/such/account"]])
    assert not Company.objects.filter(inn=loose.inn).exists()
    loose.refresh_from_db()
    assert loose.company is None


def test_upload_marks_our_company(client, finance):
    client.force_login(finance)
    upload(client, [["ООО Наша", "111222333", "да", ""], ["ООО Чужая", "444555666", "no", ""]])
    assert Company.objects.get(inn="111222333").is_own
    assert not Company.objects.get(inn="444555666").is_own


def test_upload_rejects_an_unclear_own_flag(client, finance):
    client.force_login(finance)
    upload(client, [["ООО Наша", "111222333", "может быть", ""]])
    assert not Company.objects.exists()


def test_our_companies_come_first(db):
    Company.objects.create(name="AAA", inn="1")
    Company.objects.create(name="ZZZ", inn="2", is_own=True)
    assert list(Company.objects.values_list("name", flat=True)) == ["ZZZ", "AAA"]


@pytest.fixture
def seed(tmp_path, monkeypatch):
    path = tmp_path / "companies.xlsx"
    monkeypatch.setattr(company_files, "SEED_PATH", path)
    return path


def test_project_file_downloads_or_explains_how_to_write_it(client, finance, own, accounts, seed):
    client.force_login(finance)
    url = reverse("cashbook:companies_seed")
    assert client.get(url).status_code == 302  # no file yet
    company_files.write_seed_file(seed)
    response = client.get(url)
    assert response.status_code == 200
    assert "attachment" in response["Content-Disposition"]


def test_load_project_file_button(client, finance, own, accounts, seed):
    company_files.write_seed_file(seed)
    connected, _ = accounts
    own.delete()  # the accounts stay, unconnected
    client.force_login(finance)
    client.post(reverse("cashbook:companies_import"), {"seed": "1"})
    connected.refresh_from_db()
    assert connected.company.inn == "305900839"
    assert connected.company.is_own is False  # the fixture's company was not marked as ours


def test_commands_write_and_load_a_file(own, accounts, tmp_path):
    path = tmp_path / "companies.xlsx"
    call_command("export_companies", str(path), stdout=StringIO())
    connected, _ = accounts
    own.delete()
    call_command("load_companies", str(path), stdout=StringIO())
    connected.refresh_from_db()
    assert connected.company.inn == "305900839"


def test_project_files_fill_an_empty_database(db):
    """The committed files load in order: counterparties first, then the companies that refer to them."""
    assert not load_files([exports.seed_path(kind) for kind in exports.SEED_FILES]).errors
    report = company_files.load_companies(company_files.SEED_PATH)
    assert report.errors == []
    assert Company.objects.filter(is_own=True).count() == 1
    assert Counterparty.objects.filter(company__isnull=False).count() == report.linked


def test_region_is_preferred_but_not_required(client, finance):
    tash, _ = Region.objects.get_or_create(name="Tash reg")
    client.force_login(finance)
    url = reverse("cashbook:company_add")
    assert "Tash reg" in client.get(url).content.decode()
    assert client.post(url, {"name": "С регионом", "inn": "100000001", "region": tash.pk}).status_code == 302
    assert client.post(url, {"name": "Без региона", "inn": "100000002", "region": ""}).status_code == 302
    assert Company.objects.get(inn="100000001").region == tash
    assert Company.objects.get(inn="100000002").region is None


def test_inactive_regions_are_not_offered_for_a_new_company(client, finance):
    Region.objects.create(name="Closed", is_active=False)
    client.force_login(finance)
    assert "Closed" not in client.get(reverse("cashbook:company_add")).content.decode()


def test_page_shows_the_region(client, finance, own, accounts):
    own.region, _ = Region.objects.get_or_create(name="Samarqand")
    own.save()
    client.force_login(finance)
    assert "Samarqand" in client.get(reverse("cashbook:companies")).content.decode()


def region_workbook(rows):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["Company", "INN", "Region", "Statement reference"])
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    book.save(buffer)
    return SimpleUploadedFile("companies.xlsx", buffer.getvalue())


def test_upload_sets_the_region_by_name(client, finance, own):
    Region.objects.get_or_create(name="Bukhara")
    client.force_login(finance)
    file = region_workbook([[own.name, own.inn, "bukhara", ""], ["ООО Новая", "100000003", "", ""]])
    client.post(reverse("cashbook:companies_import"), {"file": file})
    own.refresh_from_db()
    assert own.region.name == "Bukhara"
    assert Company.objects.get(inn="100000003").region is None


def test_upload_with_an_unknown_region_saves_nothing(client, finance):
    client.force_login(finance)
    file = region_workbook([["ООО Новая", "100000003", "Nowhere", ""]])
    client.post(reverse("cashbook:companies_import"), {"file": file})
    assert not Company.objects.exists()
    assert not Region.objects.filter(name="Nowhere").exists()
