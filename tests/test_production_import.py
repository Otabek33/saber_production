"""Importing products, materials, recipes and past output from the workbooks."""

import datetime
from decimal import Decimal
from io import StringIO

import openpyxl
import pytest
from django.core.management import call_command

from finance_saber.production.importer import NameMatcher, key, read_workbooks
from finance_saber.production.models import Material, MaterialMove, Product, ProductionRun, material_stock

D = Decimal


def test_key_ignores_spacing_case_and_lookalike_letters():
    assert key("Цитрат натрия Е331 ") == key("Цитрат натрия E331")  # Cyrillic and Latin E
    assert key('Основа  "барбарис"') == key("Основа барбарис")


def test_matcher_uses_aliases_and_picks_the_tray_by_product():
    matcher = NameMatcher(
        [
            "B Fresh  анор банка  0,450 л",
            "Гофра паддон saber  450 мл\\12 шт",
            "Гофра паддон saber  450 мл\\24 шт",
            "Гофра паддон sakura  450 мл\\12 шт",
        ]
    )
    assert matcher.material_for("B Fresh банка гранат 0,450 л", "B Fresh анор", D(1)) == "B Fresh  анор банка  0,450 л"
    assert matcher.material_for("Гофра поддон", "B Fresh анор", D("0.0833")).endswith("saber  450 мл\\12 шт")
    assert matcher.material_for("Гофра паттон", "7 Energy Drink", D("0.0417")).endswith("saber  450 мл\\24 шт")
    assert matcher.material_for("Гофра поддон", "Green Tea мёд", D("0.0833")).startswith("Гофра паддон sakura")
    assert matcher.material_for("Неизвестное", "B Fresh анор", D(1)) is None


@pytest.fixture
def workbooks(tmp_path):
    """A small "Сбор" (one recipe, the stock list) and "Warehouse" (two products, one day)."""
    report = openpyxl.Workbook()
    sheet = report.active
    sheet.title = "Сбор"
    sheet["B2"], sheet["C2"], sheet["E2"] = "РЕЦЕПТУРА  Green Tea мёд 450 ml", "Ед.изм.", "на 1 банку"
    sheet["B3"], sheet["E3"] = "Сырье наименование", 0.45
    lines = [
        ("Сахар", "кг", 79, 0.03555),
        ("  банка Green Tea мёд 0,450 л", "шт", 2222.2, 1),
        ("Кришка золотая 202", "шт", 2222.2, 1),
        ("Гофра паттон", "шт", 185.185, 0.08333),
        ("клей", "кг", 0.17, 7.65e-05),
    ]
    for offset, (name, unit, dose, per_unit) in enumerate(lines):
        r = 4 + offset
        sheet[f"A{r}"], sheet[f"B{r}"], sheet[f"C{r}"], sheet[f"D{r}"], sheet[f"E{r}"] = (
            offset + 1, name, unit, dose, per_unit,
        )  # fmt: skip
    sheet["W5"] = "Сырье наименование"
    stock = [
        ("Сахар", "12 месяц", 20482.15, "кг"),
        ("баночки", None, None, None),
        ("  банка Green Tea мёд 0,450 л", None, 5000, "шт "),
        ("Золотой кришка 202", None, 559404, "шт "),
        ("Гофра паддон sakura  450 мл\\12 шт", None, 18548, "шт "),
    ]
    for offset, (name, shelf, amount, unit) in enumerate(stock):
        r = 6 + offset
        sheet[f"W{r}"], sheet[f"X{r}"], sheet[f"Z{r}"], sheet[f"AA{r}"] = name, shelf, amount, unit
    report_path = tmp_path / "report.xlsx"
    report.save(report_path)

    situation = openpyxl.Workbook()
    sheet = situation.active
    sheet.title = "Warehouse"
    sheet["C1"], sheet["F1"], sheet["G1"] = "Отчёт", datetime.datetime(2026, 8, 3), datetime.datetime(2026, 8, 4)
    sheet["C2"], sheet["D2"], sheet["F2"] = "Green Tea мёд 0,450 л", "Sakura Зеленый чай с мёдом", 43848
    sheet["C3"], sheet["D3"] = "7 Energy Drink 0,250 л", "7 Energy 250 ML 1/12"
    situation_path = tmp_path / "Situation.xlsx"
    situation.save(situation_path)
    return report_path, situation_path


def test_read_workbooks(workbooks):
    books = read_workbooks(*workbooks)
    assert [m.name for m in books.materials] == [
        "Сахар", "банка Green Tea мёд 0,450 л", "Золотой кришка 202", "Гофра паддон sakura 450 мл\\12 шт",
    ]  # fmt: skip
    assert books.materials[1].unit == "шт"
    assert books.materials[1].kind == "packaging"
    [recipe] = books.recipes
    assert len(recipe.lines) == 5
    assert [p.name for p in books.products] == ["Green Tea мёд 0,450 л", "7 Energy Drink 0,250 л"]
    assert books.products[0].output == {datetime.date(2026, 8, 3): 43848}
    assert books.products[1].volume == D("0.25")


def test_import_command(db, workbooks):
    out = StringIO()
    call_command("import_production_xlsx", *map(str, workbooks), "--stock-date", "2026-10-22", stdout=out)

    tea = Product.objects.get(name="Green Tea мёд 0,450 л")
    assert {line.material.name: line.quantity_per_unit for line in tea.recipe.select_related("material")} == {
        "Сахар": D("0.0355500000"),
        "банка Green Tea мёд 0,450 л": D("1.0000000000"),
        "Золотой кришка 202": D("1.0000000000"),
        "Гофра паддон sakura 450 мл\\12 шт": D("0.0833300000"),
        "клей": D("0.0000765000"),
    }
    assert material_stock()[Material.objects.get(name="Сахар").pk] == D("20482.15")
    assert Material.objects.get(name="клей").moves.count() == 0  # not in the stock list: no stock

    run = ProductionRun.objects.get()
    assert (run.date, run.quantity, run.affects_stock) == (datetime.date(2026, 8, 3), 43848, False)
    assert not tea.moves.exists()  # history only
    assert "7 Energy Drink 0,250 л: no recipe" in out.getvalue()

    # Running again adds nothing.
    call_command("import_production_xlsx", *map(str, workbooks), "--stock-date", "2026-10-22", stdout=StringIO())
    assert MaterialMove.objects.count() == 4
    assert tea.recipe.count() == 5
    assert ProductionRun.objects.count() == 1


def test_dry_run_saves_nothing(db, workbooks):
    out = StringIO()
    call_command("import_production_xlsx", *map(str, workbooks), "--dry-run", stdout=out)
    assert not Material.objects.exists()
    assert "Dry run" in out.getvalue()


def test_norm_is_worked_out_when_the_formula_has_no_saved_value(tmp_path, workbooks):
    report_path, situation_path = workbooks
    book = openpyxl.load_workbook(report_path)
    book["Сбор"]["E4"] = None  # as saved by a program that doesn't calculate formulas
    book.save(report_path)
    [recipe] = read_workbooks(report_path, situation_path).recipes
    assert recipe.lines[0].per_unit == D("79") * D("0.45") / 1000


def upload(path):
    from django.core.files.uploadedfile import SimpleUploadedFile

    return SimpleUploadedFile(path.name, path.read_bytes())


def test_import_page_checks_then_imports(client, user, workbooks):
    client.force_login(user)
    url = "/production/import/"
    assert client.get(url).status_code == 200

    data = {"stock_date": "2026-10-22", "dry_run": "on"}
    response = client.post(url, {**data, "recipes": upload(workbooks[0]), "warehouse": upload(workbooks[1])})
    assert response.context["report"].new_products == 2
    assert not response.context["report"].saved
    assert not Product.objects.exists()

    del data["dry_run"]
    response = client.post(url, {**data, "recipes": upload(workbooks[0]), "warehouse": upload(workbooks[1])})
    assert response.context["report"].saved
    assert Product.objects.count() == 2
    assert MaterialMove.objects.filter(created_by=user).count() == 4


def test_import_page_rejects_a_file_that_is_not_a_workbook(client, user, workbooks, tmp_path):
    from django.core.files.uploadedfile import SimpleUploadedFile

    client.force_login(user)
    response = client.post(
        "/production/import/",
        {
            "stock_date": "2026-10-22",
            "recipes": SimpleUploadedFile("report.xlsx", b"not a zip"),
            "warehouse": upload(workbooks[1]),
        },
    )
    assert response.status_code == 200
    [error] = response.context["form"].non_field_errors()
    assert error.startswith("report.xlsx")  # names the file that couldn't be read
    assert not Product.objects.exists()
