"""Loading the handbooks from Excel, e.g. into a new server's empty database."""

from io import BytesIO, StringIO

import openpyxl
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import CommandError, call_command
from django.urls import reverse
from django.utils import translation

from finance_saber.cashbook import exports
from finance_saber.cashbook.handbooks import load_handbooks
from finance_saber.cashbook.models import Article, ArticleKind, CashDesk, Category, Counterparty, Region

URL = reverse("cashbook:handbook_import")
REFERENCE = "20208000900713230001/304577774/GOODWILL BUSINESS TRADE xususiy korxonasi"


def seed_file(language="ru") -> bytes:
    with translation.override(language):
        return exports.xlsx_response("seed.xlsx", exports.seed_sheets()).content


@pytest.fixture
def extra(db):
    """Records a fresh database (migrations only) doesn't have."""
    article = Article.objects.create(name="Аренда склада", kind=ArticleKind.EXPENSE, position=7, in_cash_flow=True)
    article.desks.set([CashDesk.objects.get(slug="bank")])
    Category.objects.create(article=article, name="Ташкент", is_active=False)
    region = Region.objects.create(name="Test region", manager="TDM9")
    Counterparty.objects.create(
        reference=REFERENCE, **Counterparty.parse_reference(REFERENCE), article=article, region=region
    )
    return article


def remove(extra):
    Counterparty.objects.all().delete()
    Category.objects.filter(article=extra).delete()
    extra.delete()
    Region.objects.filter(name="Test region").delete()


@pytest.mark.parametrize("language", ["ru", "en"])
def test_seed_file_fills_an_empty_database_and_loads_again_without_changes(extra, language):
    data = seed_file(language)
    remove(extra)  # as on a new server

    report = load_handbooks(BytesIO(data))
    assert report.saved and not report.errors
    article = Article.objects.get(name="Аренда склада")
    assert (article.kind, article.position, list(article.desks.values_list("slug", flat=True))) == (
        ArticleKind.EXPENSE, 7, ["bank"],
    )  # fmt: skip
    assert not article.categories.get(name="Ташкент").is_active
    counterparty = Counterparty.objects.get(reference=REFERENCE)
    assert (counterparty.inn, counterparty.article, counterparty.region.manager) == ("304577774", article, "TDM9")

    again = load_handbooks(BytesIO(data))
    assert not again.created and not again.updated
    assert again.unchanged["counterparties"] == 1


def test_bad_rows_save_nothing_and_are_named(extra):
    book = openpyxl.load_workbook(BytesIO(seed_file()))
    book["Статьи"]["C2"] = "Касса Луна"  # an unknown cash desk
    book["Контрагенты"]["D2"] = "Нет такой статьи"
    book["Контрагенты"]["L2"] = "1/2/Новый"
    buffer = BytesIO()
    book.save(buffer)

    report = load_handbooks(BytesIO(buffer.getvalue()))
    assert not report.saved
    assert report.errors == [
        "Лист «Статьи», строка 2: неизвестная касса «Касса Луна».",
        "Лист «Контрагенты», строка 2: неизвестная статья «Нет такой статьи»: сначала добавьте её в справочник.",
    ]
    assert not Counterparty.objects.filter(reference="1/2/Новый").exists()


def test_page_checks_then_loads(client, superuser, extra):
    data = seed_file()
    remove(extra)
    client.force_login(superuser)
    assert URL in client.get(reverse("cashbook:handbook")).content.decode()

    response = client.post(URL, {"file": SimpleUploadedFile("seed.xlsx", data), "dry_run": "on"})
    content = response.content.decode()
    assert "Проверено, ещё не сохранено" in content
    assert "Найдены листы: Статьи и категории, Статьи, Регионы, Контрагенты" in content
    assert not Counterparty.objects.exists()

    response = client.post(URL, {"file": SimpleUploadedFile("seed.xlsx", data)})
    assert response.context["report"].saved
    assert "Справочники загружены." in response.content.decode()
    assert Counterparty.objects.count() == 1


def test_page_rejects_files_without_handbook_sheets(client, superuser):
    client.force_login(superuser)
    book = openpyxl.Workbook()
    buffer = BytesIO()
    book.save(buffer)
    response = client.post(URL, {"file": SimpleUploadedFile("x.xlsx", buffer.getvalue())})
    assert "Листы справочников не найдены" in response.content.decode()


def test_page_needs_handbook_permissions(client, no_role):
    client.force_login(no_role)
    assert client.get(URL).status_code == 403


def test_commands(extra, tmp_path):
    path = tmp_path / "handbooks.xlsx"
    call_command("export_handbooks", str(path), stdout=StringIO())
    remove(extra)

    out = StringIO()
    call_command("load_handbooks", str(path), "--dry-run", stdout=out)
    assert "Dry run" in out.getvalue()
    assert not Counterparty.objects.exists()

    call_command("load_handbooks", str(path), stdout=StringIO())
    assert Counterparty.objects.count() == 1
    with pytest.raises(CommandError, match="File not found"):
        call_command("load_handbooks", str(tmp_path / "missing.xlsx"))


# --- The finance workbook's "Handbook" sheet, from the Counterparties page --------

PAGE = reverse("cashbook:counterparties")
UPLOAD = reverse("cashbook:counterparty_import")
SALES = "20208000900713230001/304577774/GOODWILL BUSINESS TRADE xususiy korxonasi"
SHOP = "20208000707198863005/311902477/GLOBUS ONE MCHJ"
ODD = "20208000000000000001/300000001/ODD MCHJ"


def situation_file() -> bytes:
    """A finance workbook whose "Handbook" sheet lists counterparties in H:J and Y:AA, as Situation.xlsx does."""
    book = openpyxl.Workbook()
    book.active.title = "Касса Бон"  # other sheets are ignored
    sheet = book.create_sheet("Handbook")
    sheet["H1"], sheet["Y1"] = "Наименование", "Наименование"
    sheet["H2"], sheet["I2"], sheet["J2"] = SALES, "Месный продажа", "Tash reg"
    sheet["Y2"], sheet["Z2"], sheet["AA2"] = SHOP, "Месный продажа", "Новый регион"
    sheet["H3"], sheet["I3"], sheet["J3"] = SHOP, "Месный продажа", 0  # again, region "0": nothing new
    sheet["Y3"], sheet["Z3"] = ODD, "Нет такой статьи"
    sheet["H4"], sheet["I4"] = "Итого", 12  # not a counterparty
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def messages_of(response):
    from django.contrib.messages import get_messages

    return [str(message) for message in get_messages(response.wsgi_request)]


def test_counterparties_page_loads_the_finance_workbook_handbook(client, user):
    client.force_login(user)
    assert UPLOAD in client.get(PAGE).content.decode()

    response = client.post(UPLOAD, {"file": SimpleUploadedFile("Situation.xlsx", situation_file()), "dry_run": "on"})
    assert response.status_code == 302
    assert not Counterparty.objects.exists()
    assert any(m.startswith("Проверен лист «Handbook», ничего не сохранено") for m in messages_of(response))

    response = client.post(UPLOAD, {"file": SimpleUploadedFile("Situation.xlsx", situation_file())})
    loaded = next(m for m in messages_of(response) if m.startswith("Загружено"))
    assert loaded == (
        "Загружено с листа «Handbook». Регионы: новых 1, изменено 0, без изменений 0; "
        "Контрагенты: новых 3, изменено 0, без изменений 0."
    )
    assert "Статьи «Нет такой статьи» нет в справочнике: контрагентов без статьи — 1." in messages_of(response)
    sales, shop, odd = (Counterparty.objects.get(reference=r) for r in (SALES, SHOP, ODD))
    assert (sales.article.name, sales.region.name, sales.inn) == ("Месный продажа", "Tash reg", "304577774")
    assert (shop.region.name, odd.article) == ("Новый регион", None)


def test_finance_workbook_keeps_classifications_made_in_the_app(client, user):
    supplies = Article.objects.get(name="Местные поставщики")
    Counterparty.objects.create(reference=SALES, **Counterparty.parse_reference(SALES), article=supplies)
    client.force_login(user)
    client.post(UPLOAD, {"file": SimpleUploadedFile("Situation.xlsx", situation_file())})
    sales = Counterparty.objects.get(reference=SALES)
    assert sales.article == supplies  # set in the app: kept
    assert sales.region.name == "Tash reg"  # was empty: filled from the workbook

    report = load_handbooks(BytesIO(situation_file()))
    assert not report.created and not report.updated  # loading it again changes nothing


def test_counterparties_page_shows_upload_errors(client, user):
    client.force_login(user)
    response = client.post(UPLOAD, {"file": SimpleUploadedFile("x.xlsx", b"not a workbook")})
    assert response.status_code == 302
    with translation.override("ru"):
        expected = translation.gettext("The file is not a readable .xlsx workbook.")
    assert messages_of(response) == [expected]
    assert not Counterparty.objects.exists()


# --- The project's handbook files (cashbook/seed/) ----------------------------------------


@pytest.fixture
def seed_dir(tmp_path, monkeypatch, extra):
    """The project's files written from a database that has ``extra``, in a temporary folder."""
    from finance_saber.cashbook import exports

    monkeypatch.setattr(exports, "SEED_DIR", tmp_path / "seed")
    exports.write_seed_files(tmp_path / "seed")
    return tmp_path / "seed"


def test_project_files_are_downloaded_from_both_pages(client, user, seed_dir):
    client.force_login(user)
    assert reverse("cashbook:handbook_seed") in client.get(reverse("cashbook:handbook")).content.decode()
    assert reverse("cashbook:counterparty_seed") in client.get(PAGE).content.decode()

    response = client.get(reverse("cashbook:handbook_seed"))
    assert response.status_code == 200
    assert "attachment" in response["Content-Disposition"]
    assert openpyxl.load_workbook(BytesIO(response.content)).sheetnames == ["Статьи и категории", "Статьи"]
    response = client.get(reverse("cashbook:counterparty_seed"))
    assert openpyxl.load_workbook(BytesIO(response.content)).sheetnames == ["Контрагенты", "Регионы"]


def test_project_files_are_not_public(client, no_role, seed_dir):
    assert client.get(reverse("cashbook:counterparty_seed")).status_code == 302  # to sign-in
    client.force_login(no_role)
    assert client.get(reverse("cashbook:counterparty_seed")).status_code == 403


def test_missing_project_file(client, user, tmp_path, monkeypatch):
    from finance_saber.cashbook import exports

    monkeypatch.setattr(exports, "SEED_DIR", tmp_path / "nothing")
    client.force_login(user)
    response = client.get(reverse("cashbook:handbook_seed"))
    assert response.status_code == 302
    assert any("export_handbooks --seed" in m for m in messages_of(response))


def test_project_files_fill_an_empty_database_in_one_click(client, superuser, seed_dir, extra):
    remove(extra)  # as on a new server
    client.force_login(superuser)
    page = client.get(URL).content.decode()
    assert "Загрузить файлы проекта" in page and "handbook.xlsx" in page and "counterparties.xlsx" in page

    response = client.post(URL, {"use_seed": "1", "dry_run": "on"})
    report = response.context["report"]
    assert not report.errors  # counterparties checked against the articles the handbook would add
    assert not report.saved
    assert not Counterparty.objects.exists()

    response = client.post(URL, {"use_seed": "1"})
    assert response.context["report"].saved
    counterparty = Counterparty.objects.get(reference=REFERENCE)
    assert counterparty.article.name == "Аренда склада"  # an article that came from the handbook file

    again = client.post(URL, {"use_seed": "1"}).context["report"]
    assert not again.created and not again.updated


def test_export_command_writes_the_project_files(extra, tmp_path, monkeypatch):
    from finance_saber.cashbook import exports

    monkeypatch.setattr(exports, "SEED_DIR", tmp_path / "seed")
    monkeypatch.setattr(exports.write_seed_files, "__defaults__", (tmp_path / "seed",))
    call_command("export_handbooks", "--seed", stdout=StringIO())
    assert {path.name for path in (tmp_path / "seed").iterdir()} == {"handbook.xlsx", "counterparties.xlsx"}
    with pytest.raises(CommandError, match="--seed"):
        call_command("export_handbooks", stdout=StringIO())


def test_the_project_files_load_into_the_test_database(db):
    """The files committed in cashbook/seed/ are readable and load without errors."""
    from finance_saber.cashbook import exports
    from finance_saber.cashbook.handbooks import load_files

    paths = [exports.seed_path(kind) for kind in exports.SEED_FILES]
    if not all(path.is_file() for path in paths):
        pytest.skip("the project has no handbook files")
    report = load_files(paths, dry_run=True)
    assert report.errors == []
    assert report.created["counterparties"] > 0
