"""The home dashboard and the weekly plans behind it."""

import datetime
from decimal import Decimal

import openpyxl
import pytest
from django.urls import reverse

from finance_saber.cashbook.models import (
    Account,
    Article,
    ArticleKind,
    CashDesk,
    Category,
    Counterparty,
    ExchangeRate,
    Kind,
    Operation,
    Region,
)
from finance_saber.core import dashboard
from finance_saber.planning import services
from finance_saber.planning.importer import import_workbook
from finance_saber.planning.models import CashPlan, CollectionPlan
from finance_saber.production.models import Product
from finance_saber.sales.models import Client, Plan, Sale

D = Decimal
MONDAY = datetime.date(2026, 10, 19)
PERIOD = dashboard.Period(datetime.date(2026, 10, 1), datetime.date(2026, 10, 21))


@pytest.fixture
def bon(db):
    return CashDesk.objects.get(slug="bon")


@pytest.fixture
def bank(db):
    return CashDesk.objects.get(slug="bank")


def category(article_name, kind=ArticleKind.EXPENSE, name=None):
    article, _ = Article.objects.get_or_create(name=article_name, defaults={"kind": kind})
    return Category.objects.get_or_create(article=article, name=name or article_name)[0]


def operation(desk, kind, amount, date=MONDAY, account=Account.CASH_UZS, category_=None, **extra):
    return Operation.objects.create(
        desk=desk,
        kind=kind,
        date=date,
        account=account,
        amount=D(amount),
        category=category_ or category("Питание" if kind == Kind.EXPENSE else "Продажа БОН", kind),
        **extra,
    )


# --- Period -------------------------------------------------------------------------------


def test_period_defaults_to_this_month(db):
    period = dashboard.period_from({})
    assert period.date_from.day == 1 and period.preset == "month"


def test_period_from_dates_and_swapped_dates():
    period = dashboard.period_from({"from": "2026-10-20", "to": "2026-10-01"})
    assert (period.date_from, period.date_to) == (datetime.date(2026, 10, 1), datetime.date(2026, 10, 20))


def test_week_preset_starts_on_monday():
    period = dashboard.preset_period("week", datetime.date(2026, 10, 22))
    assert period.date_from == MONDAY


# --- Money ----------------------------------------------------------------------------------


def test_cash_flow_by_article_and_account(bon, bank):
    operation(bon, Kind.INCOME, "1000")
    operation(bank, Kind.INCOME, "500", account=Account.BANK)
    operation(bon, Kind.EXPENSE, "300")
    operation(bon, Kind.INCOME, "10", account=Account.CASH_USD)
    ExchangeRate.objects.create(date=datetime.date(2026, 10, 1), rate=D("12000"))
    flow = dashboard.cash_flow(PERIOD, D("12000"))
    [income] = flow["income"]
    accounts = dict(zip(Account.values, income["accounts"], strict=True))
    assert (accounts[Account.CASH_UZS], accounts[Account.BANK], accounts[Account.CASH_USD]) == (1000, 500, 10)
    assert income["total"] == D("121500")  # 1000 + 500 + 10 USD × 12 000
    assert flow["totals"]["expense"]["total"] == 300
    assert flow["net"] == D("121200")


def test_bookkeeping_and_transfers_are_not_cash_flow(bon):
    opening = category("Начальный остаток", ArticleKind.INCOME)
    opening.article.in_cash_flow = False
    opening.article.save()
    operation(bon, Kind.INCOME, "5000", category_=opening)
    transfer = category("Переброска на карту", ArticleKind.TRANSFER)
    Operation.objects.create(
        desk=bon, kind=Kind.TRANSFER, date=MONDAY, account=Account.CASH_UZS, to_account=Account.CARD,
        amount=D("100"), to_amount=D("100"), category=transfer,
    )  # fmt: skip
    flow = dashboard.cash_flow(PERIOD, None)
    assert flow["income"] == [] and flow["expense"] == []


def test_cancelled_operations_do_not_count(bon):
    operation(bon, Kind.INCOME, "1000", is_cancelled=True)
    assert dashboard.cash_flow(PERIOD, None)["totals"]["income"]["total"] == 0


def test_accounts_opening_flows_and_closing(bon):
    operation(bon, Kind.INCOME, "1000", date=datetime.date(2026, 9, 30))
    operation(bon, Kind.INCOME, "200")
    operation(bon, Kind.EXPENSE, "50")
    row = next(r for r in dashboard.accounts(PERIOD, None) if r["desk"] == bon and r["account"] == Account.CASH_UZS)
    assert (row["opening"], row["in"], row["out"], row["closing"]) == (1000, 200, 50, 1150)


def test_exchange_rate_uses_the_latest_rate_on_or_before(db):
    ExchangeRate.objects.create(date=datetime.date(2026, 10, 1), rate=D("12000"))
    ExchangeRate.objects.create(date=datetime.date(2026, 10, 15), rate=D("12100"))
    assert dashboard.exchange_rate(datetime.date(2026, 10, 14)) == D("12000")
    assert dashboard.exchange_rate(datetime.date(2026, 10, 20)) == D("12100")


def test_daily_chart_has_a_bar_per_day(bon):
    operation(bon, Kind.INCOME, "1000")
    chart = dashboard.daily_chart(PERIOD, None)
    assert len(chart["labels"]) == 21
    assert chart["in"][18] == 1000.0


# --- Plans ----------------------------------------------------------------------------------


def test_weekly_cash_plan_against_fact(bon):
    sales_category = category("Продажа БОН", ArticleKind.INCOME)
    CashPlan.objects.create(date=MONDAY, article=sales_category.article, amount=D("800"))
    operation(bon, Kind.INCOME, "1000", category_=sales_category)
    sections, totals = services.cash_plan_grid(services.week_of(MONDAY))
    row = next(r for r in sections["income"] if r["article"] == sales_category.article)
    assert (row["cells"][0]["plan"], row["cells"][0]["fact"]) == (800, 1000)
    assert totals["income"]["days"][0] == {"plan": 800, "fact": 1000}


def test_saving_the_plan_grid(db):
    article = category("Питание").article
    days = services.week_of(MONDAY)
    services.save_cash_plan({f"plan-{article.pk}-{MONDAY.isoformat()}": "1 500 000"}, days)
    assert CashPlan.objects.get().amount == D("1500000")
    services.save_cash_plan({f"plan-{article.pk}-{MONDAY.isoformat()}": ""}, days)  # cleared: removed
    assert not CashPlan.objects.exists()
    # A day outside the week shown is ignored.
    services.save_cash_plan({f"plan-{article.pk}-2026-12-01": "5"}, days)
    assert not CashPlan.objects.exists()


def test_collections_by_region(bank):
    region = Region.objects.get_or_create(name="Samarqand")[0]  # seeded by a migration
    payer = Counterparty.objects.create(reference="1/2/SAM", name="SAM", region=region)
    CollectionPlan.objects.create(date=MONDAY, region=region, amount=D("1000"))
    operation(bank, Kind.INCOME, "600", account=Account.BANK, counterparty=payer)
    result = dashboard.collections(PERIOD)
    [row] = result["rows"]
    assert (row["plan"], row["fact"], row["percent"]) == (1000, 600, 60)


def test_month_production_and_sales_plan(db):
    product = Product.objects.create(name="Tea", volume=D("0.45"))
    Plan.objects.create(kind=Plan.Kind.SALES, year=2026, month=10, product=product, quantity=1000)
    client = Client.objects.create(name="Shop")
    Sale.objects.create(date=datetime.date(2026, 10, 5), client=client, product=product, quantity=250, amount=D("1"))
    Sale.objects.create(date=datetime.date(2026, 10, 25), client=client, product=product, quantity=99, amount=D("2"))
    result = dashboard.products_month(PERIOD, show_sales=True, show_production=False)
    [row] = result["rows"]
    assert (row["sales"]["plan"], row["sales"]["fact"], row["sales"]["percent"]) == (1000, 250, 25)
    assert result["month_progress"] == 68  # the 21st of 31 days


# --- Excel ----------------------------------------------------------------------------------


def test_import_plans_and_rates(tmp_path, db):
    book = openpyxl.Workbook()
    week = book.active
    week.title = "1-week"
    week["C1"], week["E1"] = datetime.datetime(2026, 10, 19), datetime.datetime(2026, 10, 20)
    week["C3"], week["D3"], week["E3"] = "План", "Факт", "План"
    week["B6"], week["C6"], week["E6"] = "Продажа Distr", 130000000, 115000000
    week["B28"], week["C28"] = "ИТОГО ПРИХОД", 130000000
    category("Продажа Distr", ArticleKind.INCOME)
    plan = book.create_sheet("Sells plan")
    plan["G1"], plan["G2"] = datetime.datetime(2026, 10, 19), "план"
    plan["B3"], plan["C3"], plan["G3"] = "Navoi", "TDM5", 5000000
    rates = book.create_sheet("Курс")
    for row, (day, rate) in enumerate([(1, 12000), (2, 0), (3, 23880), (4, 12010)], start=2):
        rates.cell(row, 4, datetime.datetime(2026, 5, day))
        rates.cell(row, 7, rate)
    path = tmp_path / "plans.xlsx"
    book.save(path)

    report = import_workbook(path)
    assert report.cash_plans == 2 and report.unknown == []
    assert CollectionPlan.objects.get().region.manager == "TDM5"
    assert report.new_regions == ["Navoi"]
    assert list(ExchangeRate.objects.order_by("date").values_list("rate", flat=True)) == [D("12000"), D("12010")]
    assert len(report.skipped_rates) == 1  # 23 880: a typo
    import_workbook(path)  # again: updated, not duplicated
    assert CashPlan.objects.count() == 2


# --- Pages and roles ------------------------------------------------------------------------


def test_dashboard_page_for_a_manager(client, user, bon):
    operation(bon, Kind.INCOME, "1000")
    client.force_login(user)
    response = client.get(reverse("core:home"), {"from": "2026-10-01", "to": "2026-10-21"})
    assert response.status_code == 200
    for block in ("money", "cash_flow", "chart", "sales", "debts", "products", "stock", "week", "collections"):
        assert block in response.context, block
    assert b"cash-chart-data" in response.content


def test_dashboard_shows_only_what_the_role_allows(client, planner, no_role):
    client.force_login(planner)  # Production only: stock and production, no money or sales
    context = client.get(reverse("core:home")).context
    assert "stock" in context
    assert "money" not in context and "debts" not in context and "sales" not in context
    client.force_login(no_role)
    response = client.get(reverse("core:home"))
    assert "money" not in response.context and "sales" not in response.context
    assert "ожидает назначения роли" in response.content.decode()


@pytest.mark.parametrize("name", ["planning:cash", "planning:collections", "planning:rates"])
def test_plan_pages(client, finance, name):
    client.force_login(finance)
    assert client.get(reverse(name), {"date": MONDAY.isoformat(), "edit": "1"}).status_code == 200


def test_plan_pages_save(client, finance):
    region = Region.objects.get_or_create(name="Bukhara")[0]
    client.force_login(finance)
    response = client.post(
        reverse("planning:collections") + f"?date={MONDAY.isoformat()}",
        {f"plan-{region.pk}-{MONDAY.isoformat()}": "2000000"},
    )
    assert response.status_code == 302
    assert CollectionPlan.objects.get().amount == D("2000000")
    client.force_login(finance)  # exchange rates are finance's
    client.post(reverse("planning:rates"), {"date": "2026-10-22", "rate": "12 150"})
    assert ExchangeRate.on(datetime.date(2026, 10, 23)) == D("12150")


def test_no_role_cannot_open_plans(client, no_role):
    client.force_login(no_role)
    assert client.get(reverse("planning:cash")).status_code == 403


def test_cards_split_into_cash_card_and_bank(bon, bank):
    operation(bon, Kind.INCOME, "1000")
    operation(bon, Kind.INCOME, "5", account=Account.CASH_USD)
    operation(bon, Kind.INCOME, "300", account=Account.CARD)
    operation(bank, Kind.INCOME, "2000", account=Account.BANK)
    operation(bank, Kind.EXPENSE, "700", account=Account.BANK)
    flow = dashboard.cash_flow(PERIOD, D("12000"))
    money_in = {c["key"]: (c["amount"], c["usd"]) for c in flow["totals"]["income"]["channels"]}
    assert money_in == {"cash": (D("61000"), 5), "card": (300, 0), "bank": (2000, 0)}
    net = {c["key"]: c["amount"] for c in flow["net_channels"]}
    assert net == {"cash": D("61000"), "card": 300, "bank": 1300}
    money = dashboard.money_totals(dashboard.accounts(PERIOD, D("12000")), D("12000"))
    assert {c["key"]: c["amount"] for c in money["channels"]}["bank"] == 1300


def test_a_refund_reduces_its_articles_total(bank):
    sales = category("Месный продажа", ArticleKind.BOTH)
    operation(bank, Kind.INCOME, "1000", account=Account.BANK, category_=sales)
    operation(bank, Kind.EXPENSE, "100", account=Account.BANK, category_=sales)  # money returned to a client
    flow = dashboard.cash_flow(PERIOD, None)
    [row] = flow["income"]
    assert row["total"] == 900
    assert flow["expense"] == []


def test_dashboard_cards_show_the_split(client, user, bon):
    operation(bon, Kind.INCOME, "1000")
    client.force_login(user)
    content = client.get(reverse("core:home"), {"from": "2026-10-01", "to": "2026-10-21"}).content.decode()
    assert content.count("solar:buildings-2-line-duotone") >= 4  # the bank line on each money card


def test_header_has_a_full_screen_toggle(client, user):
    client.force_login(user)
    content = client.get(reverse("core:home")).content.decode()
    assert "fullscreen-toggle" in content and 'aria-pressed="false"' in content
