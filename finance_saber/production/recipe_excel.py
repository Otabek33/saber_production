"""Recipes and materials in Excel: download them, load them back, and the project's files.

The workbook has four sheets, written by ``recipe_sheets()`` and read by
``load_recipes()``, so a downloaded file can be loaded as it is (also into a new
server's empty database):

* «Продукты» — products: name, sales name, volume, position, in use;
* «Сырьё и упаковка» — materials: name, unit, type, shelf life, minimum stock, in use;
* «Рецептуры» — one row per material of each product's recipe: the norm per
  unit (can), the same norm per 1000 l (for reading; used only when the norm per
  unit is empty), the allowed loss and the position;
* «Другие названия» — the names products go by in other systems (SmartUp), used
  to match sales and plan imports.

Records are matched by name, so loading the same file again changes nothing.
A product listed on «Рецептуры» gets exactly the lines the file gives it (lines
not in the file are removed); products not listed there keep their recipe.
Production already entered keeps the norms it was made with (RunLine). Any of
the sheets may be left out; if any row has an error, nothing is saved. Sheet
names, headers, material types and yes / no are recognised in Russian and English.

The materials file (Raw materials page) is the «Сырьё и упаковка» sheet alone,
with each material's stock in a «Остаток» column. The stock is loaded only when
asked, as an opening balance on a given date, and only for materials that have
no stock movements yet: stock is a ledger of movements, so a server already in
use is never double-counted. Loading again with corrected figures corrects that
opening balance.

The project's files (production/seed/recipes.xlsx and materials.xlsx) are
written by ``manage.py export_recipes --seed`` / ``export_materials --seed`` and
deployed with the code, so a new server gets them in one click: Recipes and
norms › Load from Excel, Raw materials › Load from Excel.
"""

from collections import Counter
from dataclasses import dataclass, field
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

from django.core.exceptions import ValidationError
from django.db import transaction
from django.http import HttpResponse
from django.utils import translation
from django.utils.translation import gettext as _
from django.utils.translation import gettext_noop

from finance_saber.cashbook.bank import find_sheet
from finance_saber.cashbook.exports import Sheet, attachment, xlsx_bytes
from finance_saber.cashbook.handbooks import LANGUAGES, HandbookError, clean, table, texts

from .models import (
    DEFAULT_ALLOWED_LOSS,
    NORM,
    PERCENT,
    QUANTITY,
    Material,
    MaterialMove,
    Product,
    ProductAlias,
    RecipeLine,
    material_stock,
)

PRODUCTS = gettext_noop("Products")
MATERIALS = gettext_noop("Materials and packaging")
RECIPES = gettext_noop("Recipes")
ALIASES = gettext_noop("Other product names")
SHEET_ORDER = (PRODUCTS, MATERIALS, RECIPES, ALIASES)  # the order they load in: recipes refer to the others

COLUMNS = {
    PRODUCTS: {
        gettext_noop("Name"): "name",
        gettext_noop("Sales name"): "sales_name",
        gettext_noop("Volume, l"): "volume",
        gettext_noop("Position"): "position",
        gettext_noop("In use"): "is_active",
    },
    MATERIALS: {
        gettext_noop("Name"): "name",
        gettext_noop("Unit"): "unit",
        gettext_noop("Type"): "kind",
        gettext_noop("Shelf life"): "shelf_life",
        gettext_noop("Minimum stock"): "min_stock",
        gettext_noop("In use"): "is_active",
        gettext_noop("Stock"): "stock",  # the materials file only; loaded as an opening balance when asked
    },
    RECIPES: {
        gettext_noop("Product"): "product",
        gettext_noop("Material"): "material",
        gettext_noop("Unit"): "unit",
        gettext_noop("Norm per unit"): "quantity_per_unit",
        gettext_noop("Norm per 1000 l"): "per_1000_l",
        gettext_noop("Allowed loss, %"): "allowed_loss",
        gettext_noop("Position"): "position",
    },
    ALIASES: {
        gettext_noop("Name"): "name",
        gettext_noop("Product"): "product",
        gettext_noop("Source"): "source",
    },
}
REQUIRED = {
    PRODUCTS: ("name",),
    MATERIALS: ("name",),
    RECIPES: ("product", "material"),
    ALIASES: ("name", "product"),
}


def headers(title: str, without=()) -> list[str]:
    return [_(header) for header, name in COLUMNS[title].items() if name not in without]


def _yes_no(value: bool) -> str:
    return _("yes") if value else _("no")


# --- Download ------------------------------------------------------------------------


def recipe_sheets() -> list[Sheet]:
    """Products, materials, recipes and other product names, as the loader reads them back."""
    products = list(Product.objects.order_by("position", "name"))
    lines = RecipeLine.objects.select_related("product", "material").order_by(
        "product__position", "product__name", "position", "id"
    )
    return [
        Sheet(
            _(PRODUCTS),
            headers(PRODUCTS),
            [
                [product.name, product.sales_name, product.volume, product.position, _yes_no(product.is_active)]
                for product in products
            ],
        ),
        material_sheet(with_stock=False),
        Sheet(
            _(RECIPES),
            headers(RECIPES),
            [
                [
                    line.product.name,
                    line.material.name,
                    line.material.unit,
                    line.quantity_per_unit.normalize(),
                    line.dose_per_1000_l.quantize(Decimal("0.0000001"), rounding=ROUND_HALF_UP).normalize(),
                    line.allowed_loss.normalize(),
                    line.position,
                ]
                for line in lines
            ],
        ),
        Sheet(
            _(ALIASES),
            headers(ALIASES),
            [
                [alias.name, alias.product.name, alias.source]
                for alias in ProductAlias.objects.select_related("product").order_by("product__position", "name")
            ],
        ),
    ]


def material_sheet(with_stock=True) -> Sheet:
    """Every material; with its current stock for the materials file."""
    stock = material_stock() if with_stock else {}
    rows = []
    for material in Material.objects.order_by("-kind", "name"):  # raw materials, then packaging
        row = [
            material.name,
            material.unit,
            Material.Kind(material.kind).label,
            material.shelf_life,
            material.min_stock,
            _yes_no(material.is_active),
        ]
        if with_stock:
            row.append(stock.get(material.pk, Decimal(0)))
        rows.append(row)
    return Sheet(_(MATERIALS), headers(MATERIALS, without=() if with_stock else ("stock",)), rows)


def material_sheets() -> list[Sheet]:
    return [material_sheet()]


# --- The project's files ---------------------------------------------------------------
#
# Kept with the code (production/seed/), so a new server with an empty database
# gets the recipes and materials from them. ``manage.py export_recipes --seed``
# and ``export_materials --seed`` write them again.

SEED_DIR = Path(__file__).resolve().parent / "seed"
SEED_FILES = {
    # kind: (file in SEED_DIR, download name, sheets)
    "recipes": ("recipes.xlsx", gettext_noop("Recipes"), recipe_sheets),
    "materials": ("materials.xlsx", gettext_noop("Materials and packaging"), material_sheets),
}


def seed_path(kind: str) -> Path:
    return SEED_DIR / SEED_FILES[kind][0]


def write_seed_file(kind: str, path: Path | None = None) -> Path:
    """Write the project's file of ``kind`` from the database (in Russian, as the company uses it)."""
    path = path or seed_path(kind)
    path.parent.mkdir(parents=True, exist_ok=True)
    with translation.override("ru"):
        path.write_bytes(xlsx_bytes(SEED_FILES[kind][2]()))
    return path


def seed_response(kind: str) -> HttpResponse | None:
    """The project's file of ``kind`` as a download, or None if the project has none."""
    path = seed_path(kind)
    if not path.is_file():
        return None
    return attachment(path.read_bytes(), f"{_(SEED_FILES[kind][1])} ({_('project file')}).xlsx")


# --- Upload --------------------------------------------------------------------------


@dataclass
class RecipeImport:
    sheets: list[str] = field(default_factory=list)  # the sheets found, as named in the file
    created: Counter = field(default_factory=Counter)  # by record type: "products", "materials", ...
    updated: Counter = field(default_factory=Counter)
    unchanged: Counter = field(default_factory=Counter)
    removed: Counter = field(default_factory=Counter)  # recipe lines the file no longer has
    errors: list[str] = field(default_factory=list)
    stock_kept: list[str] = field(default_factory=list)  # materials whose stock already has movements
    saved: bool = False

    def rows(self):
        """(label, created, updated, unchanged, removed) per record type, for the result table."""
        models = {"products": Product, "materials": Material, "recipe_lines": RecipeLine, "aliases": ProductAlias}
        labels = {key: str(model._meta.verbose_name_plural).capitalize() for key, model in models.items()}
        labels["opening_balances"] = _("Opening balances")
        return [
            (label, self.created[key], self.updated[key], self.unchanged[key], self.removed[key])
            for key, label in labels.items()
            if self.created[key] or self.updated[key] or self.unchanged[key] or self.removed[key]
        ]


class _Rollback(Exception):
    pass


def number(value, places: int, digits: int, signed: bool = False) -> Decimal | None:
    """A number from a cell: a number, or text like "0,0355", "1 250" or "1/12"; None if empty."""
    if value is None or clean(value) == "":
        return None
    if isinstance(value, bool):
        raise ValueError(_("“%(value)s” is not a number.") % {"value": value})
    if isinstance(value, float):
        result = Decimal(repr(value))
    elif isinstance(value, int | Decimal):
        result = Decimal(value)
    else:
        text = clean(value).replace(" ", "").replace(",", ".")
        try:
            if "/" in text:
                numerator, denominator = (Decimal(part) for part in text.split("/", 1))
                result = numerator / denominator
            else:
                result = Decimal(text)
        except (InvalidOperation, ArithmeticError):
            raise ValueError(_("“%(value)s” is not a number.") % {"value": clean(value)}) from None
    if not result.is_finite() or (result < 0 and not signed):
        raise ValueError(_("“%(value)s” is not a number.") % {"value": clean(value)})
    result = result.quantize(Decimal(1).scaleb(-places), rounding=ROUND_HALF_UP)
    if len(result.as_tuple().digits) > digits:
        raise ValueError(_("“%(value)s” is too large.") % {"value": clean(value)})
    return result


class RecipeLoader:
    """Upserts products, materials, recipes and other names from the sheets' rows. Call inside a transaction."""

    def __init__(self, report: RecipeImport, opening_date=None, user=None):
        self.report = report
        self.opening_date = opening_date  # load the materials' stock as opening balances on this date
        self.user = user
        self.kinds = {kind.value: kind for kind in Material.Kind}
        for language in LANGUAGES:
            with translation.override(language):
                self.kinds.update({str(kind.label).casefold(): kind for kind in Material.Kind})
        self.yes = texts(gettext_noop("yes")) | {"true", "1", "да"}
        self.no = texts(gettext_noop("no")) | {"false", "0", "нет"}
        self.products = {product.name.casefold(): product for product in Product.objects.all()}
        self.materials = {material.name.casefold(): material for material in Material.objects.all()}
        self.recipes: dict[int, tuple[Product, list]] = {}  # product id -> (product, its rows on «Рецептуры»)

    # --- values ---

    def error(self, sheet: str, row: int, message: str):
        self.report.errors.append(_("Sheet «%(sheet)s», row %(row)s: %(message)s") % {
            "sheet": sheet, "row": row, "message": message,
        })  # fmt: skip

    def flag(self, value, default=True) -> bool:
        text = clean(value).casefold()
        if not text:
            return default
        if text in self.yes:
            return True
        if text in self.no:
            return False
        raise ValueError(_("“%(value)s” is not yes or no.") % {"value": clean(value)})

    def kind(self, value, default) -> str:
        text = clean(value).casefold()
        if not text:
            return default
        kind = self.kinds.get(text)
        if kind is None:
            choices = ", ".join(str(choice.label) for choice in Material.Kind)
            raise ValueError(_("unknown type “%(value)s” (expected: %(choices)s).") % {
                "value": clean(value), "choices": choices,
            })  # fmt: skip
        return kind

    @staticmethod
    def position(value, default: int) -> int:
        result = number(value, 0, 5)
        return default if result is None else int(result)

    def product(self, value) -> Product:
        product = self.products.get(clean(value).casefold())
        if product is None:
            raise ValueError(_("unknown product “%(name)s”: add it on the «%(sheet)s» sheet.") % {
                "name": clean(value), "sheet": _(PRODUCTS),
            })  # fmt: skip
        return product

    def material(self, value) -> Material:
        material = self.materials.get(clean(value).casefold())
        if material is None:
            raise ValueError(_("unknown material “%(name)s”: add it on the «%(sheet)s» sheet.") % {
                "name": clean(value), "sheet": _(MATERIALS),
            })  # fmt: skip
        return material

    def save(self, kind: str, instance, values: dict, created: bool):
        changed = created or any(getattr(instance, name) != value for name, value in values.items())
        for name, value in values.items():
            setattr(instance, name, value)
        if changed:
            instance.full_clean()
            instance.save()
        bucket = self.report.created if created else self.report.updated if changed else self.report.unchanged
        bucket[kind] += 1

    # --- records ---

    def load_product(self, row):
        name = clean(row["name"])
        product = self.products.get(name.casefold())
        created = product is None
        if created:
            product = Product(name=name)
        values = {
            "sales_name": clean(row.get("sales_name")) if "sales_name" in row else product.sales_name,
            "volume": number(row.get("volume"), 3, 6) or product.volume,
            "position": self.position(row.get("position"), product.position),
            "is_active": self.flag(row.get("is_active"), product.is_active),
        }
        if values["volume"] is None:
            raise ValueError(_("the volume is required."))
        self.save("products", product, values, created)
        self.products[name.casefold()] = product  # only once saved: recipes may refer to it

    def load_material(self, row):
        name = clean(row["name"])
        material = self.materials.get(name.casefold())
        created = material is None
        if created:
            material = Material(name=name)
        min_stock = number(row.get("min_stock"), QUANTITY["decimal_places"], QUANTITY["max_digits"])
        values = {
            "unit": clean(row.get("unit")) or material.unit,
            "kind": self.kind(row.get("kind"), material.kind),
            "shelf_life": clean(row.get("shelf_life")) if "shelf_life" in row else material.shelf_life,
            "min_stock": min_stock if min_stock is not None else material.min_stock,
            "is_active": self.flag(row.get("is_active"), material.is_active),
        }
        if not values["unit"]:
            raise ValueError(_("the unit is required."))
        self.save("materials", material, values, created)
        self.materials[name.casefold()] = material  # only once saved: recipes may refer to it
        if self.opening_date and "stock" in row:
            self.opening_balance(material, number(row["stock"], QUANTITY["decimal_places"], 18, signed=True))

    def opening_balance(self, material: Material, stock: Decimal | None):
        """The file's stock as the material's opening balance, if it has no other stock movements."""
        source = f"seed:opening:{material.name}"
        if MaterialMove.objects.filter(material=material).exclude(source=source).exists():
            if stock:
                self.report.stock_kept.append(material.name)
            return
        move = MaterialMove.objects.filter(source=source).first()
        if not stock or stock < 0:  # nothing (or a shortage) to open with
            if move and not move.is_cancelled:
                move.is_cancelled = True
                move.save(update_fields=["is_cancelled"])
                self.report.removed["opening_balances"] += 1
            return
        if move is None:
            MaterialMove.objects.create(
                material=material,
                date=self.opening_date,
                kind=MaterialMove.Kind.OPENING,
                quantity=stock,
                note=_("Opening balance from Excel"),
                source=source,
                created_by=self.user,
            )
            self.report.created["opening_balances"] += 1
        elif (move.quantity, move.date, move.is_cancelled) != (stock, self.opening_date, False):
            move.quantity, move.date, move.is_cancelled = stock, self.opening_date, False
            move.save(update_fields=["quantity", "date", "is_cancelled"])
            self.report.updated["opening_balances"] += 1
        else:
            self.report.unchanged["opening_balances"] += 1

    def collect_recipe_line(self, number_, row):
        """Recipe rows are saved per product (see save_recipes), after every row is checked."""
        product = self.product(row["product"])
        self.recipes.setdefault(product.pk, (product, []))[1].append((number_, row))

    def save_recipes(self, sheet: str):
        for product, rows in self.recipes.values():
            existing = {line.material_id: line for line in product.recipe.all()}
            seen = set()
            for position, (row_number, row) in enumerate(rows, start=1):
                try:
                    material = self.material(row["material"])
                    if material.pk in seen:
                        raise ValueError(_("“%(name)s” is already in this product's recipe.") % {"name": material})
                    seen.add(material.pk)
                    per_unit = number(row.get("quantity_per_unit"), NORM["decimal_places"], NORM["max_digits"])
                    if per_unit is None:  # only the norm per 1000 l is given: the same per can of this volume
                        per_1000_l = number(row.get("per_1000_l"), NORM["decimal_places"], NORM["max_digits"])
                        if per_1000_l is not None:
                            per_unit = (per_1000_l * product.volume / 1000).quantize(
                                Decimal(1).scaleb(-NORM["decimal_places"]), rounding=ROUND_HALF_UP
                            )
                    if not per_unit:
                        raise ValueError(_("the norm is required and must be above zero."))
                    allowed = number(row.get("allowed_loss"), PERCENT["decimal_places"], PERCENT["max_digits"])
                    line = existing.get(material.pk)
                    created = line is None
                    if created:
                        line = RecipeLine(product=product, material=material)
                    values = {
                        "quantity_per_unit": per_unit,
                        "allowed_loss": DEFAULT_ALLOWED_LOSS if allowed is None else allowed,
                        "position": self.position(row.get("position"), position),
                    }
                    self.save("recipe_lines", line, values, created)
                except (ValueError, ValidationError) as error:
                    message = "; ".join(error.messages) if isinstance(error, ValidationError) else str(error)
                    self.error(sheet, row_number, message)
            removed = [line.pk for material_id, line in existing.items() if material_id not in seen]
            if removed:
                RecipeLine.objects.filter(pk__in=removed).delete()
                self.report.removed["recipe_lines"] += len(removed)

    def load_alias(self, row):
        name = clean(row["name"])
        product = self.product(row["product"])
        alias = ProductAlias.objects.filter(name__iexact=name).first()
        created = alias is None
        if created:
            alias = ProductAlias(name=name)
        values = {"product": product, "source": clean(row.get("source")) if "source" in row else alias.source}
        self.save("aliases", alias, values, created)


def load_recipes(file, dry_run: bool = False, only=SHEET_ORDER, opening_date=None, user=None) -> RecipeImport:
    """Load whichever of the sheets in ``only`` the workbook has; nothing is saved on errors or a dry run.

    With ``opening_date``, the materials' «Остаток» becomes their opening balance on that date.
    """
    import openpyxl

    try:
        workbook = openpyxl.load_workbook(file, data_only=True, read_only=True)
    except Exception as error:  # zip / xml errors: not a readable .xlsx
        raise HandbookError(_("The file is not a readable .xlsx workbook.")) from error
    report = RecipeImport()
    try:
        sheets, tables = {}, {}
        for title in only:
            for text in texts(title):
                found = find_sheet(workbook.sheetnames, text)
                if found:
                    sheets[title] = found
                    tables[title] = table(workbook[found], COLUMNS[title], REQUIRED[title])
                    report.sheets.append(found)
                    break
        if not sheets:
            names = ", ".join(f"«{_(title)}»" for title in only)
            raise HandbookError(_("No recipe sheets found. Expected any of: %(names)s.") % {"names": names})
        report.sheets.sort(key=workbook.sheetnames.index)  # as the file has them
    finally:
        workbook.close()

    try:
        with transaction.atomic():
            loader = RecipeLoader(report, opening_date=opening_date, user=user)
            steps = {
                PRODUCTS: lambda number_, row: loader.load_product(row),
                MATERIALS: lambda number_, row: loader.load_material(row),
                RECIPES: loader.collect_recipe_line,
                ALIASES: lambda number_, row: loader.load_alias(row),
            }
            for title in SHEET_ORDER:
                for row_number, row in tables.get(title, []):
                    try:
                        steps[title](row_number, row)
                    except (ValueError, ValidationError) as error:
                        message = "; ".join(error.messages) if isinstance(error, ValidationError) else str(error)
                        loader.error(sheets[title], row_number, message)
                if title == RECIPES and RECIPES in sheets:
                    loader.save_recipes(sheets[RECIPES])
            if report.errors or dry_run:
                raise _Rollback
            report.saved = True
    except _Rollback:
        pass
    return report
