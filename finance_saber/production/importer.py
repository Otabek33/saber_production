"""Read products, materials, recipes and production history from the Excel workbooks.

Two workbooks feed production:

* the report workbook's "Сбор" sheet: one recipe block per product (norm per can in
  column E) and, on the right (columns W–AA), the list of materials with their
  stock (column Z);
* the finance workbook's "Warehouse" sheet: the product list (column C, sales name
  in D) with units produced per day (dates in row 1 from column F).

The sheets name the same material differently ("B Fresh банка гранат 0,450 л" in a
recipe is "B Fresh анор банка 0,450 л" in the stock list), so names are matched on
a normalized key, then through ALIASES; trays depend on the product.

The workbooks come either as files on disk (the import_production_xlsx command) or
as files uploaded on the Import page; both are read in memory, so a server needs
no access to anyone's computer or shared folder. Formula cells are read as the
values Excel saved with them.
"""

import datetime
import re
from dataclasses import dataclass, field
from decimal import Decimal

from django.db import transaction
from django.utils.translation import gettext as _
from openpyxl import load_workbook

from .models import Material, MaterialMove, Product, ProductionRun, RecipeLine

RECIPE_SHEET = "Сбор"
WAREHOUSE_SHEET = "Warehouse"

# Recipe name -> stock-list name, for names that differ beyond spacing and spelling of the letters.
ALIASES = {
    "B Fresh банка гранат 0,450 л": "B Fresh анор банка 0,450 л",
    "B Fresh банка клубника 0,450 л": "B Fresh клубника банка 0,450 л",
    "B Fresh банка лайм 0,450 л": "B Fresh лайм банка 0,450 л",
    "банка B Fresh апельсин 0,450 л": "B Fresh апельсин банка 0,450 л",
    "банка B Fresh олма 0,450 л": "B Fresh олма банка 0,450 л",
    "Energy Drink банка 0,450 л": "7 Energy Drink банка 0,450 л",
    "банка Energy Pomegranat 0,450 л": "7 Energy Pomegranat банка0,450 л",
    "7 Energy Drink Mango банка 0,450 л": "7 Energy Drink mango банка 0,450 л",
    "7 Energy Drink цитрус банка": "7 Energy Drink апелсин банка 0,450 л",
    "Кришка золотая 202": "Золотой кришка 202",
    # The recipe has code 354, the stock list only 353: taken to be the same flavour.
    "Ар. Манго KG-318-354-8": "Ар. Манго KG 318-353-8",
    # D07717 in the recipe, D07177 in the stock list.
    "Ар.натуралный лимон D07717": "Ар.натуралный лимон D07177",
    "плёнка 36 см": "плёнка",
}

# Consumables the recipes use that the stock list doesn't have: created with no stock.
EXTRA_MATERIALS = {"клей": "кг", "плёнка": "кг", "стрич": "кг", "агрегационный код": "шт"}

TRAY_NAMES = ("Гофра поддон", "Гофра паттон")
TRAY_SABER_12 = "Гофра паддон saber 450 мл\\12 шт"
TRAY_SABER_24 = "Гофра паддон saber 450 мл\\24 шт"
TRAY_SAKURA_12 = "Гофра паддон sakura 450 мл\\12 шт"

PACKAGING_WORDS = ("банка", "кришка", "гофра", "клей", "плёнка", "стрич", "агрегац")

# Warehouse product -> title of its recipe block in "Сбор" (both normalized when compared).
RECIPE_FOR_PRODUCT = {
    "B Fresh анор 0,450 л": "РЕЦЕПТУРА B Fresh гранат",
    "B Fresh клубника 0,450 л": "РЕЦЕПТУРА B Fresh клубника",
    "B Fresh лимон 0,450 л": "РЕЦЕПТУРА B Fresh лайм",
    "B Fresh Манго 0,450 л": "РЕЦЕПТУРА B Fresh Mango",
    "B Fresh апельсин 0,450 л": "РЕЦЕПТУРА B Fresh апельсин",
    "B Fresh олма 0,450 л": "РЕЦЕПТУРА B Fresh олма",
    "7 Energy Drink 0,450 л": "РЕЦЕПТУРА Energy Drink 450мл",
    "7 Energy Drink Pomegranat 0,450 л": "7 Energy Drink Pomegranat 0,450 л Hyper power",
    "7 Energy Drink mango 0,450 л": "7 Energy Drink Mango 0,450 л Reactive Power",
    "7 Energy Drink цитрус 0,450 л": "7 Energy Drink цитрус 0,450 л Dinamik Power",
    "Green Tea мёд 0,450 л": "РЕЦЕПТУРА Green Tea мёд 450 ml",
    "Green Tea манго - ромашка0,450 л": "РЕЦЕПТУРА Green Tea манго - ромашка 450 ml",
    "Green Tea малина 0,450 л": "РЕЦЕПТУРА Green Tea малина 450 ml",
    "Blak Tea Яблока - корица-гвоздика 0,450 л": "РЕЦЕПТУРА Blak Tea Яблока - корица-гвоздика 450 ml",
    "Blak Tea Апельсин-корица 0,450 л": "РЕЦЕПТУРА Blak Tea Апельсин-корица 450 ml",
    "Blak Tea лимон-мёд-имбир 0,450 л": "РЕЦЕПТУРА Blak Tea лимон-мёд-имбир 450 ml",
}

# Cyrillic letters that look like Latin ones; the sheets mix them ("Е331" and "E331").
_LOOKALIKES = str.maketrans("аеорсухкмтнвіё", "aeopcyxkmthbie")


class SheetError(Exception):
    pass


def clean(name) -> str:
    """The name as stored: surrounding spaces removed, inner runs of spaces collapsed."""
    return re.sub(r"\s+", " ", str(name or "")).strip()


def key(name) -> str:
    """Matching key: lower case, look-alike letters unified, only letters and digits kept."""
    return re.sub(r"[^0-9a-zа-я]", "", clean(name).lower().translate(_LOOKALIKES))


def to_decimal(value) -> Decimal:
    if value in (None, ""):
        return Decimal(0)
    if isinstance(value, str):
        value = value.replace(" ", "").replace(",", ".")
    return Decimal(str(value))


@dataclass
class MaterialRow:
    name: str
    unit: str
    shelf_life: str
    stock: Decimal

    @property
    def kind(self) -> str:
        return "packaging" if any(word in self.name.lower() for word in PACKAGING_WORDS) else "raw"


@dataclass
class RecipeRow:
    material: str  # the name as written in the recipe
    unit: str
    per_1000_l: Decimal
    per_unit: Decimal


@dataclass
class RecipeBlock:
    title: str
    volume: Decimal
    lines: list[RecipeRow] = field(default_factory=list)


@dataclass
class ProductRow:
    name: str
    sales_name: str
    position: int
    output: dict[datetime.date, int]  # units produced per day

    @property
    def volume(self) -> Decimal:
        match = re.search(r"0,(\d+)\s*л", self.name)
        return Decimal(f"0.{match.group(1)}") if match else Decimal("0.45")


def _sheet(file, name: str):
    """All rows of sheet ``name`` as value tuples; ``file`` is a path or an uploaded file."""
    label = getattr(file, "name", str(file))
    try:
        workbook = load_workbook(file, data_only=True, read_only=True)
    except Exception as error:  # zip / xml errors: not a readable .xlsx
        raise SheetError(_("%(file)s is not a readable .xlsx workbook.") % {"file": label}) from error
    try:  # read_only keeps the file open until closed
        if name not in workbook.sheetnames:
            raise SheetError(
                _("%(file)s has no sheet “%(sheet)s” (sheets: %(sheets)s).")
                % {"file": label, "sheet": name, "sheets": ", ".join(workbook.sheetnames)}
            )
        return list(workbook[name].iter_rows(values_only=True))
    finally:
        workbook.close()


def _cell(rows, row: int, column: int):
    """1-based like Excel; empty for cells beyond the sheet's used range."""
    line = rows[row - 1] if row - 1 < len(rows) else ()
    return line[column - 1] if column - 1 < len(line) else None


def read_materials(rows) -> list[MaterialRow]:
    """The stock list on the right of "Сбор": № V, name W, shelf life X, stock Z, unit AA."""
    header = next((r for r in range(1, len(rows) + 1) if clean(_cell(rows, r, 23)) == "Сырье наименование"), None)
    if header is None:
        raise SheetError(_("The Сбор sheet has no materials list (column W, “Сырье наименование”)."))
    materials = []
    for r in range(header + 1, len(rows) + 1):
        name, unit = clean(_cell(rows, r, 23)), clean(_cell(rows, r, 27))
        if name and unit:  # rows without a unit are section headings ("баночки")
            materials.append(MaterialRow(name, unit.lower(), clean(_cell(rows, r, 24)), to_decimal(_cell(rows, r, 26))))
    return materials


def read_recipes(rows) -> list[RecipeBlock]:
    """Recipe blocks: a title row (C "Ед.изм."), a row with the can volume in E, then numbered lines."""
    starts = [r for r in range(1, len(rows) + 1) if clean(_cell(rows, r, 3)).lower().startswith("ед.изм")]
    blocks = []
    for index, start in enumerate(starts):
        end = starts[index + 1] if index + 1 < len(starts) else len(rows) + 1
        block = RecipeBlock(clean(_cell(rows, start, 2)), to_decimal(_cell(rows, start + 1, 5)))
        for r in range(start + 2, end):
            name, dose, per_unit = clean(_cell(rows, r, 2)), _cell(rows, r, 4), _cell(rows, r, 5)
            if not name or not isinstance(dose, int | float) or dose <= 0:
                continue
            if isinstance(per_unit, int | float) and per_unit > 0:
                per_unit = to_decimal(per_unit)
            elif block.volume:
                # E is a formula (dose × can volume / 1000); a file saved by a program that doesn't
                # calculate formulas has no value for it, so work it out the same way.
                per_unit = to_decimal(dose) * block.volume / 1000
            else:
                continue
            block.lines.append(RecipeRow(name, clean(_cell(rows, r, 3)).lower(), to_decimal(dose), per_unit))
        blocks.append(block)
    return blocks


def read_products(rows) -> list[ProductRow]:
    """Warehouse: product C, sales name D, units per day from F with the dates in row 1."""
    dates = {}
    for column in range(6, len(rows[0]) + 1 if rows else 0):
        value = _cell(rows, 1, column)
        if isinstance(value, datetime.datetime | datetime.date):
            dates[column] = value.date() if isinstance(value, datetime.datetime) else value
    products = []
    for r in range(2, len(rows) + 1):
        name = clean(_cell(rows, r, 3))
        if not name:
            continue
        output = {}
        for column, date in dates.items():
            value = _cell(rows, r, column)
            if isinstance(value, int | float) and value > 0:
                output[date] = int(round(value))
        products.append(ProductRow(name, clean(_cell(rows, r, 4)), len(products) + 1, output))
    return products


@dataclass
class Workbooks:
    materials: list[MaterialRow]
    recipes: list[RecipeBlock]
    products: list[ProductRow]


def read_workbooks(recipes_file, warehouse_file) -> Workbooks:
    """Read both workbooks; each is a path or an uploaded file."""
    recipe_rows = _sheet(recipes_file, RECIPE_SHEET)
    return Workbooks(
        materials=read_materials(recipe_rows),
        recipes=read_recipes(recipe_rows),
        products=read_products(_sheet(warehouse_file, WAREHOUSE_SHEET)),
    )


class NameMatcher:
    """Find the stock-list material a recipe line means."""

    def __init__(self, material_names):
        self.by_key = {key(name): name for name in material_names}
        self.aliases = {key(recipe): key(stock) for recipe, stock in ALIASES.items()}

    def material_for(self, recipe_name: str, product_name: str, per_unit: Decimal) -> str | None:
        if key(recipe_name) in {key(tray) for tray in TRAY_NAMES}:
            return self.tray_for(product_name, per_unit)
        wanted = key(recipe_name)
        wanted = self.aliases.get(wanted, wanted)
        return self.by_key.get(wanted)

    def tray_for(self, product_name: str, per_unit: Decimal) -> str | None:
        """Teas go in Sakura trays, the rest in Saber ones; about 1/24 per can means a 24-can tray."""
        if "tea" in product_name.lower():
            name = TRAY_SAKURA_12
        elif per_unit < Decimal("0.06"):
            name = TRAY_SABER_24
        else:
            name = TRAY_SABER_12
        return self.by_key.get(key(name))


# Saving to the database --------------------------------------------------------------------

NORM_STEP = Decimal("0.0000000001")


class DryRun(Exception):
    """Raised inside the import transaction to roll it back after a check."""


@dataclass
class ImportReport:
    materials: int = 0
    new_materials: int = 0
    opening_balances: int = 0
    products: int = 0
    new_products: int = 0
    recipes: int = 0
    history_days: int = 0
    without_recipe: list[str] = field(default_factory=list)  # products with no recipe block
    unmatched: list[str] = field(default_factory=list)  # "product: recipe material" not in the stock list
    saved: bool = False


def import_workbooks(books: Workbooks, stock_date: datetime.date, source: str, user=None, dry_run=False):
    """Create what's missing: materials with opening stock, products, recipes, past output.

    Existing materials, products, recipes and opening balances are left alone, so
    importing the same workbooks again changes nothing. Past output is history
    only: the stock in Сбор already reflects it. Nothing is saved when a recipe
    material can't be matched, or for a dry run; the report says what would happen.
    """
    report = ImportReport()
    try:
        with transaction.atomic():
            materials = _import_materials(books.materials, stock_date, source, user, report)
            _import_products(books, materials, user, report)
            if dry_run or report.unmatched:
                raise DryRun
            report.saved = True
    except DryRun:
        pass
    return report


def _import_materials(rows, stock_date, source, user, report) -> dict[str, Material]:
    """{matching key: Material}; new materials get an opening balance from the sheet's stock."""
    keys = [key(row.name) for row in rows]
    duplicates = sorted({row.name for row in rows if keys.count(key(row.name)) > 1})
    if duplicates:
        raise SheetError(
            _("The stock list has materials that differ only in spelling: %(names)s") % {"names": "; ".join(duplicates)}
        )
    result = {}
    for row in rows:
        material, is_new = Material.objects.get_or_create(
            name=row.name, defaults={"unit": row.unit, "kind": row.kind, "shelf_life": row.shelf_life}
        )
        report.new_materials += is_new
        result[key(row.name)] = material
        if row.stock:
            _move, is_new = MaterialMove.objects.get_or_create(
                source=f"{source}:opening:{row.name}",
                defaults={
                    "material": material,
                    "date": stock_date,
                    "kind": MaterialMove.Kind.OPENING,
                    "quantity": row.stock,
                    "note": "Остаток из Excel",
                    "created_by": user,
                },
            )
            report.opening_balances += is_new
    for name, unit in EXTRA_MATERIALS.items():
        if key(name) not in result:
            material, is_new = Material.objects.get_or_create(
                name=name, defaults={"unit": unit, "kind": Material.Kind.PACKAGING}
            )
            report.new_materials += is_new
            result[key(name)] = material
    report.materials = len(result)
    return result


def _import_products(books, materials, user, report):
    blocks = {key(block.title): block for block in books.recipes}
    recipe_titles = {key(product): key(title) for product, title in RECIPE_FOR_PRODUCT.items()}
    matcher = NameMatcher(material.name for material in materials.values())
    report.products = len(books.products)

    for row in books.products:
        product, is_new = Product.objects.get_or_create(
            name=row.name,
            defaults={"sales_name": row.sales_name, "volume": row.volume, "position": row.position * 10},
        )
        report.new_products += is_new
        block = blocks.get(recipe_titles.get(key(row.name), ""))
        if block is None:
            report.without_recipe.append(row.name)
        elif not product.recipe.exists():
            lines = {}
            for position, line in enumerate(block.lines, start=1):
                name = matcher.material_for(line.material, row.name, line.per_unit)
                if name is None:
                    report.unmatched.append(f"{row.name}: {line.material}")
                    continue
                material = materials[key(name)]
                norm = line.per_unit.quantize(NORM_STEP)
                if material.pk in lines:  # the same material twice: add up
                    lines[material.pk].quantity_per_unit += norm
                else:
                    lines[material.pk] = RecipeLine(
                        product=product, material=material, quantity_per_unit=norm, position=position
                    )
            RecipeLine.objects.bulk_create(lines.values())
            report.recipes += 1

        for date, quantity in row.output.items():
            _run, is_new = ProductionRun.objects.get_or_create(
                date=date,
                product=product,
                defaults={
                    "quantity": quantity,
                    "affects_stock": False,
                    "note": "Из Excel (Warehouse)",
                    "created_by": user,
                },
            )
            report.history_days += is_new
