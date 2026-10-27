"""Load the plans and exchange rates kept in the finance workbook.

* "1-week": day columns (dates in row 1, "План" under them in row 3), one row per
  cash-flow article (name in B, or in A when B is empty) → CashPlan.
* "Sells plan": day columns (dates in row 1, "план" in row 2), one row per sales
  region (name in B, its manager's code in C) → CollectionPlan and Region.manager.
* "Курс": dates in D and the day's rate in G (blank or 0 on days without a new
  rate) → ExchangeRate. A rate far off its neighbours is a typo and is skipped.
"""

import datetime
from dataclasses import dataclass, field
from decimal import Decimal

from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook

from finance_saber.cashbook.models import Article, ExchangeRate, Region
from finance_saber.production.importer import clean, key

from .models import CashPlan, CollectionPlan

RATE_JUMP = Decimal("0.25")  # a day's rate more than 25% away from the last good one is a typo
TOTAL = key("Итого")


@dataclass
class PlanImportReport:
    cash_plans: int = 0
    collection_plans: int = 0
    managers: int = 0
    new_regions: list[str] = field(default_factory=list)
    rates: int = 0
    skipped_rates: list[str] = field(default_factory=list)
    unknown: list[str] = field(default_factory=list)


def _date(value):
    if isinstance(value, datetime.datetime):
        return value.date()
    return value if isinstance(value, datetime.date) else None


def _amount(value) -> Decimal:
    return Decimal(str(value)) if isinstance(value, int | float) and value > 0 else Decimal(0)


def _plan_columns(rows, date_row, label_row, label):
    """{column: date} for the plan columns: a date in ``date_row`` and ``label`` under it."""
    columns = {}
    for column, value in enumerate(rows[date_row]):
        day = _date(value)
        under = rows[label_row][column] if column < len(rows[label_row]) else None
        if day and key(under) == key(label):
            columns[column] = day
    return columns


def read_cash_plan(rows, report, user=None):
    articles = {key(article.name): article for article in Article.objects.all()}
    columns = _plan_columns(rows, 0, 2, "План")
    for values in rows[3:]:
        label = clean(values[1] if len(values) > 1 else "") or clean(values[0] if values else "")
        if not label or key(label).startswith(TOTAL):
            continue
        plans = {day: _amount(values[column]) for column, day in columns.items() if column < len(values)}
        plans = {day: amount for day, amount in plans.items() if amount}
        if not plans:
            continue
        article = articles.get(key(label))
        if article is None:
            report.unknown.append(f"1-week: {label}")
            continue
        for day, amount in plans.items():
            CashPlan.objects.update_or_create(
                date=day, article=article, defaults={"amount": amount, "updated_by": user}
            )
            report.cash_plans += 1


def read_collection_plan(rows, report, user=None):
    regions = {key(region.name): region for region in Region.objects.all()}
    columns = _plan_columns(rows, 0, 1, "план")
    for values in rows[2:]:
        name = clean(values[1] if len(values) > 1 else "")
        if not name or key(name).startswith(TOTAL):
            continue
        region = regions.get(key(name))
        if region is None:  # a sales region with no bank counterparties yet
            region = regions[key(name)] = Region.objects.create(name=name)
            report.new_regions.append(name)
        manager = clean(values[2] if len(values) > 2 else "")
        if manager and not region.manager:
            region.manager = manager
            region.save(update_fields=["manager"])
            report.managers += 1
        for column, day in columns.items():
            amount = _amount(values[column]) if column < len(values) else Decimal(0)
            if amount:
                CollectionPlan.objects.update_or_create(
                    date=day, region=region, defaults={"amount": amount, "updated_by": user}
                )
                report.collection_plans += 1


def read_rates(rows, report):
    last, today = None, timezone.localdate()
    for values in rows[1:]:
        day = _date(values[3] if len(values) > 3 else None)
        rate = _amount(values[6] if len(values) > 6 else None)
        if not day or not rate or day > today:  # the sheet runs on for years of placeholder rows
            continue
        if last and abs(rate - last) / last > RATE_JUMP:
            report.skipped_rates.append(f"{day:%d.%m.%Y}: {rate} (previous {last})")
            continue
        ExchangeRate.objects.update_or_create(date=day, defaults={"rate": rate})
        report.rates += 1
        last = rate


SHEETS = {"1-week": read_cash_plan, "Sells plan": read_collection_plan}


@transaction.atomic
def import_workbook(path, user=None) -> PlanImportReport:
    report = PlanImportReport()
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        for name, reader in SHEETS.items():
            if name in workbook.sheetnames:
                reader(list(workbook[name].iter_rows(values_only=True)), report, user=user)
        if "Курс" in workbook.sheetnames:
            read_rates(list(workbook["Курс"].iter_rows(values_only=True)), report)
    finally:
        workbook.close()
    return report
