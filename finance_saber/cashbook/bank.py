"""Bank statements: read an .xlsx statement and record its lines on a bank desk.

A statement has one line per payment: date and time, the counterparty as
"account/INN/name" (Счет/ИНН), document number, operation code, МФО, Оборот
Дебет (money out) or Оборот Кредит (money in), and the payment purpose. Columns
are found by their header text, so both the bank's own export and the
workbook's "Банк" sheet can be read.

Each line becomes an Operation on the bank desk. Its counterparty is looked up
(or created) by the statement reference; the counterparty's article classifies
the line. Lines of counterparties without an article, or whose article doesn't
fit the line's direction, go to "Не распределено" until the counterparty is
classified (see reclassify()). A line already recorded is recognised by its
date and time, document number, counterparty account and amount (and, for lines
identical in all of these, their order in the file), so uploading an
overlapping statement does not duplicate anything.
"""

import datetime
import re
from collections import Counter
from dataclasses import dataclass, field
from decimal import ROUND_HALF_UP, Decimal

from django.db import transaction
from django.utils.translation import gettext as _

from .models import Account, Article, CashDesk, Category, Counterparty, Kind, Operation

# Shared articles for lines that can't be classified yet (also used by the cash sheet import).
UNSORTED = {
    Kind.INCOME: "Не распределено (приход)",
    Kind.EXPENSE: "Не распределено (расход)",
}
UNSORTED_CATEGORY = "Требует проверки"

# Header text (lower case, Cyrillic) that identifies each statement column.
HEADERS = {
    "reference": ("счет/инн", "счёт/инн"),
    "document": ("№ док", "номер док"),
    "operation": ("оп",),
    "mfo": ("мфо",),
    "debit": ("оборот дебет", "дебет"),
    "credit": ("оборот кредит", "кредит"),
    "purpose": ("назначение платежа", "назначение"),
}
REQUIRED = ("date", "reference", "debit", "credit")
WORKBOOK_SHEET = "Банк"  # the statement sheet of the finance workbook (Situation.xlsx)
MAX_HEADER_ROW = 20
CENT = Decimal("0.01")
# Latin letters that bank exports mix into Cyrillic headers ("Cчет/ИНН").
LOOKALIKES = str.maketrans("caeopxyCAEOPXY", "саеорхуСАЕОРХУ")


class StatementError(Exception):
    """The file isn't a bank statement this reader understands."""


@dataclass
class StatementLine:
    row: int
    posted_at: datetime.datetime
    reference: str
    debit: Decimal | None = None
    credit: Decimal | None = None
    document: str = ""
    operation: str = ""
    mfo: str = ""
    purpose: str = ""
    # 1 for the first line with this content in the file, 2 for an identical second one, ...
    # Two separate payments can match in every column; this keeps both.
    occurrence: int = 1

    @property
    def kind(self) -> str:
        return Kind.INCOME if self.credit else Kind.EXPENSE

    @property
    def amount(self) -> Decimal:
        return self.credit or self.debit

    @property
    def key(self) -> str:
        """Identifies the line across statements, to skip ones already recorded."""
        account = Counterparty.parse_reference(self.reference)["account_number"] or self.reference[:40]
        key = f"bank:{self.posted_at:%Y%m%d%H%M%S}/{self.document}/{account}/{self.kind}/{self.amount}"
        return key if self.occurrence == 1 else f"{key}#{self.occurrence}"


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", " ", str(value).replace("\xa0", " ")).strip()


def _money(value) -> Decimal | None:
    if isinstance(value, bool) or not isinstance(value, int | float | Decimal):
        value = _text(value).replace(" ", "").replace(",", ".")
        if not re.fullmatch(r"-?\d+(\.\d+)?", value):
            return None
    amount = Decimal(str(value)).quantize(CENT, rounding=ROUND_HALF_UP)
    return amount if amount > 0 else None


def find_sheet(sheetnames: list[str], wanted: str) -> str | None:
    """The workbook's sheet named like ``wanted``, ignoring case and spaces ("Касса бон ", "КАССА БОН")."""

    def key(name):
        return re.sub(r"\s+", "", name).casefold()

    return next((name for name in sheetnames if key(name) == key(wanted)), None)


def sheet_not_found(wanted: str, sheetnames: list[str]) -> str:
    return _("Sheet «%(name)s» not found. The workbook has: %(sheets)s.") % {
        "name": wanted,
        "sheets": ", ".join(f"«{name}»" for name in sheetnames),
    }


def _header(value) -> str:
    return _text(value).translate(LOOKALIKES).lower()


def _find_columns(rows) -> tuple[int, dict[str, int], list[int]]:
    """The header row's index, the column of each field, and every "Дата" column."""
    for index, row in enumerate(rows[:MAX_HEADER_ROW]):
        headers = [_header(cell) for cell in row]
        columns, dates = {}, []
        for position, text in enumerate(headers):
            if text.startswith("дата"):
                dates.append(position)
            for name, names in HEADERS.items():
                if name not in columns and any(text == n or text.startswith(n) for n in names):
                    columns[name] = position
        if dates:
            columns["date"] = dates[0]
        if all(name in columns for name in REQUIRED):
            return index, columns, dates
    raise StatementError(_("Couldn't find the statement header (columns Дата, Счет/ИНН, Оборот Дебет, Оборот Кредит)."))


def read_statement(file, sheet_name: str | None = None) -> list[StatementLine]:
    """Statement lines from an .xlsx file (a path or an uploaded file), in file order."""
    import openpyxl

    try:
        workbook = openpyxl.load_workbook(file, data_only=True, read_only=True)
    except Exception as error:  # zip / xml errors: not a readable .xlsx
        raise StatementError(_("The file is not a readable .xlsx workbook.")) from error
    try:
        if sheet_name:
            found = find_sheet(workbook.sheetnames, sheet_name)
            if found is None:
                raise StatementError(sheet_not_found(sheet_name, workbook.sheetnames))
        else:
            # The finance workbook keeps its statement lines on "Банк"; a bank export has one sheet.
            found = find_sheet(workbook.sheetnames, WORKBOOK_SHEET)
        sheet = workbook[found] if found else workbook.worksheets[0]
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()

    header_index, columns, date_columns = _find_columns(rows)

    def value(row, name):
        position = columns.get(name)
        return row[position] if position is not None and position < len(row) else None

    lines = []
    for number, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        fields = {name: value(row, name) for name in columns}
        # Several "Дата" columns (e.g. a date-only helper column): prefer one with a time.
        dates = [row[c] for c in date_columns if c < len(row) and isinstance(row[c], datetime.datetime)]
        posted_at = max(dates, key=lambda d: d.time() != datetime.time(0), default=None)
        reference = _text(fields.get("reference"))
        debit, credit = _money(fields.get("debit")), _money(fields.get("credit"))
        if not posted_at or not reference or not (debit or credit):
            continue  # totals, blank and zero lines
        if debit and credit:
            raise StatementError(_("Row %(row)s has both a debit and a credit amount.") % {"row": number})
        lines.append(
            StatementLine(
                row=number,
                posted_at=posted_at,
                reference=reference,
                debit=debit,
                credit=credit,
                document=_text(fields.get("document")),
                operation=_text(fields.get("operation")),
                mfo=_text(fields.get("mfo")),
                purpose=_text(fields.get("purpose")),
            )
        )
    seen = Counter()
    for line in lines:
        seen[line.key] += 1
        line.occurrence = seen[line.key]
    return lines


@dataclass
class BankImportReport:
    created: int = 0
    duplicates: int = 0
    new_counterparties: list[Counterparty] = field(default_factory=list)
    unclassified: int = 0  # lines filed under "Не распределено"
    created_dates: set[datetime.date] = field(default_factory=set)
    # With create_counterparties=False: references of counterparties not in the app yet, and
    # how many of their lines were left out (they are recorded by a later upload).
    unknown_references: list[str] = field(default_factory=list)
    skipped_unknown: int = 0


def unsorted_category(kind: str) -> Category:
    article, _ = Article.objects.get_or_create(name=UNSORTED[kind], defaults={"kind": kind, "position": 999})
    category, _ = Category.objects.get_or_create(article=article, name=UNSORTED_CATEGORY)
    return category


def bank_category(desk: CashDesk, counterparty: Counterparty, kind: str) -> tuple[Category, bool]:
    """The category for a bank line, and whether it is classified (not "Не распределено")."""
    article = counterparty.article
    if article and article.is_active and article.accepts(kind) and article.is_available_on(desk):
        category, _ = Category.objects.get_or_create(article=article, name=article.name)
        return category, True
    return unsorted_category(kind), False


@transaction.atomic
def record_lines(
    desk: CashDesk, lines: list[StatementLine], user=None, create_counterparties: bool = True
) -> BankImportReport:
    """Record statement lines on a bank desk, skipping lines recorded before.

    A line's counterparty is created when it isn't known yet; with
    ``create_counterparties=False`` such lines are left out instead, so the
    counterparties can be added by hand first (with their article).
    """
    report = BankImportReport()
    existing = set(desk.operations.filter(source__startswith="bank:").values_list("source", flat=True))
    counterparties = {c.reference: c for c in Counterparty.objects.select_related("article")}
    for line in lines:
        if line.key in existing:
            report.duplicates += 1
            continue
        counterparty = counterparties.get(line.reference)
        if counterparty is None and not create_counterparties:
            report.skipped_unknown += 1
            if line.reference not in report.unknown_references:
                report.unknown_references.append(line.reference)
            continue
        if counterparty is None:
            counterparty = Counterparty.objects.create(
                reference=line.reference, **Counterparty.parse_reference(line.reference)
            )
            counterparties[line.reference] = counterparty
            report.new_counterparties.append(counterparty)
        category, classified = bank_category(desk, counterparty, line.kind)
        operation = Operation(
            desk=desk,
            kind=line.kind,
            date=line.posted_at.date(),
            category=category,
            account=Account.BANK,
            amount=line.amount,
            counterparty=counterparty,
            document_number=line.document[:50],
            note=line.purpose[:500],
            source=line.key,
            created_by=user,
        )
        operation.full_clean(validate_unique=False)
        operation.save()
        existing.add(line.key)
        report.created += 1
        report.created_dates.add(operation.date)
        report.unclassified += not classified
    return report


def reclassify(counterparty: Counterparty) -> int:
    """Re-file the counterparty's bank lines under its current article. Returns how many changed."""
    changed = 0
    for operation in counterparty.operations.select_related("desk", "category"):
        if not operation.desk.is_bank:
            continue
        category, _ = bank_category(operation.desk, counterparty, operation.kind)
        if operation.category_id != category.pk:
            operation.category = category
            operation.save(update_fields=["category", "updated_at"])
            changed += 1
    return changed


def unclassified_counterparties():
    """Counterparties that still need an article."""
    return Counterparty.objects.filter(article__isnull=True)
