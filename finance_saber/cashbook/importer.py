"""Import a cash sheet of the finance workbook (Situation.xlsx) into a cash desk.

Two sheet layouts are known (see LAYOUTS):

* "Касса Бон": up to six amounts per row, Приход UZS / USD / Карта (columns H, I,
  J) and Расход UZS / USD / Карта (L, M, N), plus the cash-flow article (B),
  category (C), organisation (D), date (E), exchange rate (F) and note (G).
* "Инкасса & Розница": UZS only, Приход (F) and Расход (G), with the article (B),
  date (D) and note (E), and an opening balance in H5. Its "Дист расход" rows
  get a category from the note (стоянка, арава, клик, ...).

Each row becomes one or more Operations:

* a single amount                          -> one income or expense
* one Приход + one Расход, other currency  -> a transfer with the amount received (обмен валюты)
* one Приход + one Расход, same currency,
  equal amounts                            -> a transfer between the accounts (cash <-> card)
* anything else                            -> one income or expense per amount
* a negative amount                        -> the opposite direction (e.g. a refunded expense)

The sheet's running balance (Остаток, columns P, Q, R) is not always the sum of
its rows: it can be typed in by hand, e.g. after a cash count (in Situation.xlsx
it only starts on row 242). Wherever the sheet's balance differs from the running
total of the imported rows, a "Корректировка остатка" operation is added for the
difference, so the imported history ends on the sheet's balances and each
correction is visible (and can be cancelled) in the app.

Categories are matched to the handbook by name; a known article with a new
category name gets that category added; rows without a usable article go to a
"Не распределено" article for later review.

Each operation remembers its sheet row (Operation.source), so the workbook can be
imported again and again, e.g. uploaded every day while it is still kept in
Excel: rows imported before are skipped. A skipped row is compared with the
operation it became; if its date or amounts changed in Excel, or rows were
inserted or deleted above it, the import reports a conflict instead of guessing.
"""

import copy
import datetime
import re
from collections import Counter
from collections.abc import Callable
from dataclasses import dataclass, field
from decimal import ROUND_HALF_UP, Decimal

from django.db import transaction
from django.utils.translation import gettext as _

# Articles that collect rows the sheet left unclassified (shared with bank statements).
from .bank import UNSORTED, UNSORTED_CATEGORY, find_sheet, sheet_not_found
from .ledger import balances
from .models import ACCOUNT_CURRENCY, BOOKKEEPING_ARTICLES, Account, Article, CashDesk, Category, Kind, Operation

COLUMN_LETTERS = "ABCDEFGHIJKLMNOPQR"

# Articles for differences between the sheet's typed-in balance and the sum of its rows.
ADJUSTMENT = {
    Kind.INCOME: "Корректировка остатка (+)",
    Kind.EXPENSE: "Корректировка остатка (−)",
}
ADJUSTMENT_CATEGORY = "Сверка с Excel"
# Differences up to this are float rounding in the sheet, not real gaps.
BALANCE_TOLERANCE = Decimal("0.05")
# Handbook categories used when a row is recognised as a transfer but its article isn't one.
TRANSFER_CATEGORY = ("Переброска на карту", "Переброска на карту")
EXCHANGE_CATEGORY = ("Обмена валюты", "Обмена валюты")

CENT = Decimal("0.01")
OPENING_BALANCE = ("Начальный остаток", "Начальный остаток")

# "Инкасса & Розница": category of a "Дист расход" row from words in its note, first match wins
# (cartridges before fuel: "заправка картриджа" is not fuel).
DIST_EXPENSE_CATEGORIES = [
    ("Стоянка", ("стоян",)),
    ("Арава", ("арава",)),
    ("Клик", ("клик", "слик")),
    ("Брак", ("брак",)),
    ("Питание", ("обед", "ужен", "ужин")),
    ("Карз и долг", ("карз", "долг")),
    ("Картридж и печать", ("картр", "катр", "котр", "распечат")),
    ("Топливо", ("газ", "бензин", "солярк", "заправ")),
    ("Машина", ("вулкан", "машин", "ремонт", "гаи")),
    ("Корректировка", ("коррект",)),
]


def inkassa_category(article: str, note: str) -> str:
    """Category for a "Инкасса & Розница" row: from the note for "Дист расход", else the article itself."""
    if article != "Дист расход":
        return article
    lowered = note.lower()
    for category, words in DIST_EXPENSE_CATEGORIES:
        if any(word in lowered for word in words):
            return category
    return "Прочие"


@dataclass(frozen=True)
class SheetLayout:
    """Where a sheet keeps each value: column letters, and the header row's expected text."""

    date: str
    note: str
    income: dict[str, str]  # column letter -> account
    expense: dict[str, str]
    balance: dict[str, str]
    expected_headers: dict[str, str]
    article: str = ""
    category: str = ""
    organization: str = ""
    rate: str = ""
    desk: str = "bon"  # default cash desk slug for this sheet
    header_row: int = 4
    first_data_row: int = 6
    opening_balance_row: int | None = None  # a row holding only the balance the sheet starts from
    category_from_note: Callable[[str, str], str] | None = None  # (article, note) -> category


LAYOUTS = {
    "Касса Бон": SheetLayout(
        article="B",
        category="C",
        organization="D",
        date="E",
        rate="F",
        note="G",
        income={"H": Account.CASH_UZS, "I": Account.CASH_USD, "J": Account.CARD},
        expense={"L": Account.CASH_UZS, "M": Account.CASH_USD, "N": Account.CARD},
        balance={"P": Account.CASH_UZS, "Q": Account.CASH_USD, "R": Account.CARD},
        expected_headers={"C": "Категория", "E": "Дата", "G": "Примичание", "H": "UZS", "I": "USD", "J": "Карта"},
        desk="bon",
    ),
    "Инкасса & Розница": SheetLayout(
        article="B",
        organization="C",
        date="D",
        note="E",
        income={"F": Account.CASH_UZS},
        expense={"G": Account.CASH_UZS},
        balance={"H": Account.CASH_UZS},
        expected_headers={"B": "Категория", "D": "Дата", "E": "Примичание", "F": "UZS", "G": "UZS"},
        desk="inkassa",
        opening_balance_row=5,
        category_from_note=inkassa_category,
    ),
}


# Cash desk slug -> the sheet kept for it in the finance workbook.
SHEETS_BY_DESK = {layout.desk: name for name, layout in LAYOUTS.items()}

# Sources name the workbook the history was first imported from, whatever the
# uploaded copy is called ("Situation (2).xlsx"), so every upload continues it.
SOURCE_WORKBOOK = "Situation.xlsx"


def source_prefix(sheet_name: str) -> str:
    return f"{SOURCE_WORKBOOK}/{sheet_name}"


class SheetError(Exception):
    """The workbook doesn't have the expected layout."""


def clean_name(value) -> str:
    """Handbook names as typed in the sheet: trim, drop non-breaking and repeated spaces."""
    if not isinstance(value, str):
        return ""
    return re.sub(r"\s+", " ", value.replace("\xa0", " ")).strip()


def to_money(value) -> Decimal | None:
    """A sheet number rounded to cents (sheet formulas leave float noise), or None for blanks and zero."""
    if isinstance(value, bool) or not isinstance(value, int | float | Decimal):
        return None
    amount = Decimal(str(value)).quantize(CENT, rounding=ROUND_HALF_UP)
    return amount or None


@dataclass
class SheetRow:
    number: int
    date: datetime.date
    article: str = ""
    category: str = ""
    organization: str = ""
    note: str = ""
    rate: Decimal | None = None
    income: dict[str, Decimal] = field(default_factory=dict)
    expense: dict[str, Decimal] = field(default_factory=dict)
    balance: dict[str, Decimal] = field(default_factory=dict)  # the sheet's Остаток after this row, where filled in
    date_filled: bool = False  # the sheet had no date; the previous row's date was used

    @property
    def has_amounts(self) -> bool:
        return bool(self.income or self.expense)

    @property
    def full_note(self) -> str:
        note = self.note
        if self.organization and self.organization.lower() not in note.lower():
            note = f"{note} · {self.organization}" if note else self.organization
        return note[:255]


@dataclass
class PlannedOperation:
    kind: str
    account: str
    amount: Decimal
    to_account: str = ""
    to_amount: Decimal | None = None
    # Which handbook category to use: the row's own, or a fixed one for recognised transfers.
    category: tuple[str, str] | None = None
    is_adjustment: bool = False

    def movements(self) -> dict[str, Decimal]:
        if self.kind == Kind.INCOME:
            return {self.account: self.amount}
        if self.kind == Kind.EXPENSE:
            return {self.account: -self.amount}
        return {self.account: -self.amount, self.to_account: self.to_amount}


def plan_row(row: SheetRow) -> list[PlannedOperation]:
    """The operations one sheet row stands for (see the module docstring)."""
    income, expense = dict(row.income), dict(row.expense)
    # A negative amount is a movement in the other direction (e.g. "Бонус возврат").
    for account, amount in list(income.items()):
        if amount < 0:
            del income[account]
            expense[account] = expense.get(account, 0) - amount
    for account, amount in list(expense.items()):
        if amount < 0:
            del expense[account]
            income[account] = income.get(account, 0) - amount

    if len(income) == 1 and len(expense) == 1:
        (to_account, received), (from_account, sent) = next(iter(income.items())), next(iter(expense.items()))
        if from_account != to_account:
            if ACCOUNT_CURRENCY[from_account] != ACCOUNT_CURRENCY[to_account]:
                return [PlannedOperation(Kind.TRANSFER, from_account, sent, to_account, received, EXCHANGE_CATEGORY)]
            if sent == received:
                return [PlannedOperation(Kind.TRANSFER, from_account, sent, to_account, received, TRANSFER_CATEGORY)]

    planned = [PlannedOperation(Kind.INCOME, account, amount) for account, amount in income.items()]
    planned += [PlannedOperation(Kind.EXPENSE, account, amount) for account, amount in expense.items()]
    return planned


@dataclass
class Sheet:
    rows: list[SheetRow]
    closing: dict[str, Decimal]  # the sheet's own final Остаток, for checking the import
    opening: dict[str, Decimal] = field(default_factory=dict)  # the balance the sheet starts from


def read_sheet(path, sheet_name: str, layout: SheetLayout | None = None) -> Sheet:
    """The sheet's rows with amounts, its opening balance and its final balances.

    ``path`` is a file path or an uploaded file. Reading stops at the last row
    with a date; below it the sheet holds loose calculations, not entries. A row
    with amounts but no date takes the date of the row before it.
    """
    layout = layout or LAYOUTS.get(sheet_name)
    if layout is None:
        raise SheetError(f"No layout for sheet {sheet_name!r}. Known sheets: {', '.join(LAYOUTS)}")
    import openpyxl  # only needed for imports; see the "excel" dependency group

    try:
        workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except OSError:
        raise
    except Exception as error:  # zip / xml errors: not a readable .xlsx
        raise SheetError(_("The file is not a readable .xlsx workbook.")) from error
    found = find_sheet(workbook.sheetnames, sheet_name)
    if found is None:
        sheetnames = workbook.sheetnames
        workbook.close()
        raise SheetError(sheet_not_found(sheet_name, sheetnames))
    sheet = workbook[found]

    letters = COLUMN_LETTERS
    index = {letter: position for position, letter in enumerate(letters)}
    first_row = min(layout.header_row, layout.opening_balance_row or layout.header_row)
    # Read-only rows stop at the last filled cell; pad them so every column letter has a value.
    all_rows = [
        tuple(values) + (None,) * (len(letters) - len(values))
        for values in sheet.iter_rows(min_row=first_row, max_col=len(letters), values_only=True)
    ]
    workbook.close()

    def row_at(number):
        position = number - first_row
        return all_rows[position] if 0 <= position < len(all_rows) else (None,) * len(letters)

    header = row_at(layout.header_row)
    for letter, expected in layout.expected_headers.items():
        if clean_name(header[index[letter]]) != expected:
            raise SheetError(
                f"Unexpected sheet layout: cell {letter}{layout.header_row} is {header[index[letter]]!r}, "
                f"expected {expected!r}."
            )

    def balances_in(cell):
        return {
            account: to_money(cell[letter]) or Decimal(0)
            for letter, account in layout.balance.items()
            if isinstance(cell[letter], int | float) and not isinstance(cell[letter], bool)
        }

    opening = {}
    if layout.opening_balance_row:
        opening = balances_in(dict(zip(letters, row_at(layout.opening_balance_row), strict=True)))

    data = [(n, row_at(n)) for n in range(layout.first_data_row, first_row + len(all_rows))]
    date_index = index[layout.date]
    last_dated = max((n for n, values in data if isinstance(values[date_index], datetime.datetime)), default=0)

    rows, closing, previous_date = [], dict(opening), None
    for number, values in data:
        if number > last_dated:
            break
        cell = dict(zip(letters, values, strict=True))
        income = {account: amount for letter, account in layout.income.items() if (amount := to_money(cell[letter]))}
        expense = {account: amount for letter, account in layout.expense.items() if (amount := to_money(cell[letter]))}
        balance = balances_in(cell)
        closing.update(balance)
        if not income and not expense and not balance:
            continue
        date = cell[layout.date]
        date_filled = not isinstance(date, datetime.datetime)
        if date_filled:
            if previous_date is None:
                if income or expense:
                    raise SheetError(f"Row {number} has amounts but no date, and no earlier row has one.")
                continue
            date = previous_date
        else:
            date = date.date()
        previous_date = date
        article, note = clean_name(cell.get(layout.article)), clean_name(cell.get(layout.note))
        category = clean_name(cell.get(layout.category))
        if layout.category_from_note and article:
            category = layout.category_from_note(article, note)
        rate = to_money(cell[layout.rate]) if layout.rate else None
        rows.append(
            SheetRow(
                number=number,
                date=date,
                article=article,
                category=category,
                organization=clean_name(cell.get(layout.organization)),
                note=note,
                rate=rate if rate and rate > 0 else None,
                income=income,
                expense=expense,
                balance=balance,
                date_filled=date_filled,
            )
        )
    return Sheet(rows, closing, opening)


@dataclass(frozen=True)
class Movement:
    """What an operation does to the balances, for comparing a sheet row with what it was imported as."""

    date: datetime.date
    kind: str
    account: str
    amount: Decimal
    to_account: str = ""
    to_amount: Decimal | None = None

    @classmethod
    def of(cls, operation: Operation) -> "Movement":
        return cls(
            operation.date,
            operation.kind,
            operation.account,
            operation.amount,
            operation.to_account,
            operation.to_amount,
        )


@dataclass(frozen=True)
class Conflict:
    """A sheet row imported before whose date or amounts are different now."""

    row: int
    operation: Operation  # what the row was imported as
    sheet: Movement  # what the row says now


@dataclass
class ImportReport:
    rows: int = 0
    created: Counter = field(default_factory=Counter)  # by kind
    created_dates: set[datetime.date] = field(default_factory=set)
    skipped_existing: int = 0
    unsorted_rows: list[int] = field(default_factory=list)
    new_categories: set[str] = field(default_factory=set)
    filled_dates: list[int] = field(default_factory=list)
    reclassified_transfers: list[int] = field(default_factory=list)
    adjustments: list[tuple[int, str, Decimal]] = field(default_factory=list)  # (row, account, difference)
    opening: dict[str, Decimal] = field(default_factory=dict)
    conflicts: list[Conflict] = field(default_factory=list)
    updated: list[Conflict] = field(default_factory=list)  # with update_changed: rows rewritten from the sheet


class CashSheetImporter:
    """Creates Operations for sheet rows on one cash desk. Call inside a transaction.

    With ``update_changed``, a row imported before whose date or amounts changed in
    Excel rewrites its operation from the sheet (Excel being the official record)
    instead of being reported as a conflict.
    """

    def __init__(self, desk: CashDesk, source_prefix: str, user=None, update_changed: bool = False):
        self.desk = desk
        self.source_prefix = source_prefix
        self.user = user
        self.update_changed = update_changed
        self.report = ImportReport()
        self.articles = {article.name: article for article in Article.objects.all()}
        self.categories = {(c.article.name, c.name): c for c in Category.objects.select_related("article")}
        self.existing = {
            operation.source: operation for operation in desk.operations.filter(source__startswith=f"{source_prefix}/")
        }

    def import_rows(self, rows: list[SheetRow], opening: dict[str, Decimal] | None = None) -> ImportReport:
        running = {account: Decimal(0) for account in Account.values}
        self.report.opening = dict(opening or {})
        if opening and rows:
            # The balance the sheet starts from, as an income on the first day.
            start = SheetRow(number=0, date=rows[0].date, note="Начальный остаток из Excel")
            for account, amount in opening.items():
                if amount < 0:
                    raise SheetError(f"Negative opening balance for {account}: {amount}")
                if amount > 0:
                    planned = PlannedOperation(Kind.INCOME, account, amount, category=OPENING_BALANCE)
                    self.create(start, f"opening-{account}", planned)
                    running[account] += amount
        for row in rows:
            if row.has_amounts:
                self.report.rows += 1
                if row.date_filled:
                    self.report.filled_dates.append(row.number)
            for part, planned in enumerate(plan_row(row), start=1):
                self.create(row, part, planned)
                for account, amount in planned.movements().items():
                    running[account] += amount
            # Bring the running total to the sheet's own balance where they differ.
            for account, sheet_balance in row.balance.items():
                difference = sheet_balance - running[account]
                if abs(difference) > BALANCE_TOLERANCE:
                    kind = Kind.INCOME if difference > 0 else Kind.EXPENSE
                    adjustment = PlannedOperation(kind, account, abs(difference), is_adjustment=True)
                    if self.create(row, f"adj-{account}", adjustment):
                        self.report.adjustments.append((row.number, account, difference))
                    running[account] = sheet_balance
        return self.report

    def create(self, row: SheetRow, part, planned: PlannedOperation) -> bool:
        """Create the operation for one part of a row; False if it was imported before."""
        source = f"{self.source_prefix}/{row.number or 'opening'}/{part}"
        if source in self.existing:
            self.report.skipped_existing += 1
            operation = self.existing[source]
            sheet = Movement(
                row.date, planned.kind, planned.account, planned.amount, planned.to_account, planned.to_amount
            )
            if Movement.of(operation) != sheet:
                conflict = Conflict(row.number, copy.copy(operation), sheet)  # as it was before any update
                if self.update_changed:
                    self.fill(operation, row, planned)
                    operation.full_clean(validate_unique=False)
                    operation.save()
                    self.report.updated.append(conflict)
                else:
                    self.report.conflicts.append(conflict)
            return False
        operation = Operation(desk=self.desk, source=source, created_by=self.user)
        self.fill(operation, row, planned)
        operation.full_clean(validate_unique=False)
        operation.save()
        self.report.created[planned.kind] += 1
        self.report.created_dates.add(row.date)
        return True

    def fill(self, operation: Operation, row: SheetRow, planned: PlannedOperation) -> None:
        """Set what the sheet row says on ``operation`` (a new one, or one being updated)."""
        operation.kind = planned.kind
        operation.date = row.date
        operation.category = self.category_for(row, planned)
        operation.account = planned.account
        operation.amount = planned.amount
        operation.to_account = planned.to_account
        operation.to_amount = planned.to_amount
        operation.rate = row.rate
        operation.note = row.full_note
        if planned.is_adjustment:
            operation.note = (
                f"Корректировка до остатка в Excel (строка {row.number}): {row.balance[planned.account]:,.2f}"
            )

    def category_for(self, row: SheetRow, planned: PlannedOperation) -> Category:
        if planned.is_adjustment:
            return self.category(self.article(ADJUSTMENT[planned.kind], planned.kind), ADJUSTMENT_CATEGORY)
        if planned.category == OPENING_BALANCE:
            return self.category(self.article(OPENING_BALANCE[0], Kind.INCOME), OPENING_BALANCE[1])
        article = self.articles.get(row.article)
        if article and article.accepts(planned.kind):
            if row.category:
                return self.category(article, row.category)
            # Known article, no category: keep the article, mark the row for review.
            self.report.unsorted_rows.append(row.number)
            return self.category(article, UNSORTED_CATEGORY)
        if planned.kind == Kind.TRANSFER:
            if not (article and article.kind == Kind.TRANSFER):
                self.report.reclassified_transfers.append(row.number)
            article_name, category_name = planned.category
            if article and article.kind == Kind.TRANSFER and row.category:
                article_name, category_name = article.name, row.category
            return self.category(self.article(article_name, Kind.TRANSFER), category_name)
        self.report.unsorted_rows.append(row.number)
        return self.category(self.article(UNSORTED[planned.kind], planned.kind), UNSORTED_CATEGORY)

    def article(self, name: str, kind: str) -> Article:
        if name not in self.articles:
            self.articles[name] = Article.objects.create(
                name=name, kind=kind, position=999, in_cash_flow=name not in BOOKKEEPING_ARTICLES
            )
        return self.articles[name]

    def category(self, article: Article, name: str) -> Category:
        key = (article.name, name)
        if key not in self.categories:
            self.categories[key] = Category.objects.create(article=article, name=name)
            self.report.new_categories.add(f"{article.name} › {name}")
        return self.categories[key]


class _Rollback(Exception):
    """Undoes an upload that is only checked, or that has conflicts."""


@dataclass
class SheetUpload:
    report: ImportReport
    sheet_closing: dict[str, Decimal]  # the sheet's own final Остаток
    imported_closing: dict[str, Decimal]  # the same accounts summed over the imported operations
    saved: bool
    last_date: datetime.date | None = None  # the sheet's latest row with amounts

    @property
    def balance_differences(self) -> dict[str, tuple[Decimal, Decimal]]:
        """Accounts whose balance in the app differs from the sheet: account -> (app, sheet)."""
        return {
            account: (self.imported_closing[account], expected)
            for account, expected in self.sheet_closing.items()
            if abs(self.imported_closing[account] - expected) > BALANCE_TOLERANCE
        }


def upload_sheet(desk: CashDesk, file, user=None, dry_run: bool = False, update_changed: bool = False) -> SheetUpload:
    """Import the desk's sheet from an uploaded finance workbook; only rows not imported before are added.

    Nothing is saved for a dry run, or when rows imported before changed in Excel
    (see SheetUpload.report.conflicts) unless ``update_changed`` rewrites them from the sheet.
    """
    sheet_name = SHEETS_BY_DESK.get(desk.slug)
    if sheet_name is None:
        raise SheetError(f"Cash desk {desk.slug!r} has no sheet in the finance workbook.")
    parsed = read_sheet(file, sheet_name)
    prefix = source_prefix(sheet_name)
    try:
        with transaction.atomic():
            importer = CashSheetImporter(desk, prefix, user=user, update_changed=update_changed)
            report = importer.import_rows(parsed.rows, opening=parsed.opening)
            imported = balances(desk.operations.filter(source__startswith=f"{prefix}/"))
            upload = SheetUpload(
                report,
                parsed.closing,
                imported,
                saved=not (dry_run or report.conflicts),
                last_date=max((row.date for row in parsed.rows if row.has_amounts), default=None),
            )
            if not upload.saved:
                raise _Rollback
    except _Rollback:
        pass
    return upload
