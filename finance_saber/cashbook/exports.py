"""Excel (.xlsx) downloads of the handbook and the counterparties list.

Each sheet is a plain table: a bold, frozen header row with filters, then one
row per record, so it can be sorted and filtered in Excel straight away.
"""

import datetime
from collections.abc import Iterable, Sequence
from dataclasses import dataclass
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from django.db.models import Count, DecimalField, Prefetch, Q, QuerySet, Sum, Value
from django.db.models.functions import Coalesce
from django.http import HttpResponse
from django.utils import timezone, translation
from django.utils.functional import Promise
from django.utils.http import content_disposition_header
from django.utils.translation import gettext as _
from django.utils.translation import gettext_noop

from .models import Article, ArticleKind, Category, Counterparty, Kind, Region

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MONEY_FORMAT = "#,##0.00"
DATE_FORMAT = "DD.MM.YYYY"
MAX_COLUMN_WIDTH = 60


@dataclass
class Sheet:
    title: str
    headers: Sequence[str]
    rows: Iterable[Sequence]
    money_columns: Sequence[int] = ()  # 0-based columns shown as amounts


def xlsx_response(filename: str, sheets: Sequence[Sheet]) -> HttpResponse:
    """An .xlsx download with one table per sheet."""
    return attachment(xlsx_bytes(sheets), filename)


def attachment(content: bytes, filename: str) -> HttpResponse:
    response = HttpResponse(content, content_type=XLSX_CONTENT_TYPE)
    response["Content-Disposition"] = content_disposition_header(as_attachment=True, filename=filename)
    return response


def xlsx_bytes(sheets: Sequence[Sheet]) -> bytes:
    """An .xlsx workbook with one table per sheet."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="E8EEF9")
    for spec in sheets:
        sheet = workbook.create_sheet(spec.title[:31])  # Excel's limit for sheet names
        sheet.append(list(spec.headers))
        for cell in sheet[1]:
            cell.font, cell.fill = header_font, header_fill
        widths = [len(str(header)) for header in spec.headers]
        for values in spec.rows:
            sheet.append([str(value) if isinstance(value, Promise) else value for value in values])
            row = sheet[sheet.max_row]
            for position, cell in enumerate(row):
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.data_type = "s"  # text, never a formula Excel would run
                if isinstance(cell.value, datetime.date):
                    cell.number_format = DATE_FORMAT
                if position in spec.money_columns:
                    cell.number_format = MONEY_FORMAT
                widths[position] = max(widths[position], len(_display(cell.value)))
        for position, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(position)].width = min(width + 2, MAX_COLUMN_WIDTH)
        sheet.freeze_panes = "A2"
        if sheet.max_row > 1:
            sheet.auto_filter.ref = sheet.dimensions
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _display(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float | int | Decimal) and not isinstance(value, bool):
        return f"{value:,.2f}"
    return str(value)


def _yes_no(value: bool) -> str:
    return _("yes") if value else _("no")


def dated_filename(name: str) -> str:
    return f"{name} {timezone.localdate().isoformat()}.xlsx"


# --- Handbook --------------------------------------------------------------------


def handbook_sheets() -> list[Sheet]:
    """Articles with their categories (one row each), and the articles on their own."""
    articles = list(
        Article.objects.prefetch_related(
            "desks", Prefetch("categories", queryset=Category.objects.order_by("name"))
        ).order_by("kind", "position", "name")
    )

    def desks(article):
        return ", ".join(desk.name for desk in article.desks.all()) or _("All cash desks")

    flat = []
    for article in articles:
        common = [ArticleKind(article.kind).label, article.name]
        categories = list(article.categories.all())
        for category in categories or [None]:
            flat.append(
                [
                    *common,
                    category.name if category else "",
                    desks(article),
                    _yes_no(article.in_cash_flow),
                    _yes_no(article.is_active),
                    _yes_no(category.is_active) if category else "",
                ]
            )
    return [
        Sheet(
            _("Articles and categories"),
            [
                _("Type"),
                _("Article"),
                _("Category"),
                _("Cash desks"),
                _("Counts in cash flow"),
                _("Article active"),
                _("Category active"),
            ],
            flat,
        ),
        Sheet(
            _("Articles"),
            [
                _("Type"),
                _("Article"),
                _("Cash desks"),
                _("Position"),
                _("Counts in cash flow"),
                _("Active"),
                _("Categories"),
            ],  # fmt: skip
            [
                [
                    ArticleKind(article.kind).label,
                    article.name,
                    desks(article),
                    article.position,
                    _yes_no(article.in_cash_flow),
                    _yes_no(article.is_active),
                    len(article.categories.all()),
                ]
                for article in articles
            ],
        ),
    ]


# --- Counterparties --------------------------------------------------------------

ZERO = Value(0, output_field=DecimalField(max_digits=18, decimal_places=2))


def counterparty_sheet(counterparties: QuerySet[Counterparty]) -> Sheet:
    """The (already filtered) counterparties, with their bank lines and money in / out."""
    posted = Q(operations__is_cancelled=False)
    counterparties = counterparties.select_related("article", "region").annotate(
        line_count=Count("operations", filter=posted),
        money_in=Coalesce(Sum("operations__amount", filter=posted & Q(operations__kind=Kind.INCOME)), ZERO),
        money_out=Coalesce(Sum("operations__amount", filter=posted & Q(operations__kind=Kind.EXPENSE)), ZERO),
    )
    rows = [
        [
            counterparty.name,
            counterparty.inn,
            counterparty.account_number,
            counterparty.article.name if counterparty.article else _("Needs an article"),
            ArticleKind(counterparty.article.kind).label if counterparty.article else "",
            counterparty.region.name if counterparty.region else "",
            counterparty.region.manager if counterparty.region else "",
            counterparty.line_count,
            counterparty.money_in,
            counterparty.money_out,
            timezone.localtime(counterparty.created_at).date(),
            counterparty.reference,
        ]
        for counterparty in counterparties
    ]
    headers = [
        _("Name"),
        _("INN"),
        _("Account number"),
        _("Article"),
        _("Type"),
        _("Region"),
        _("Sales manager"),
        _("Bank lines"),
        _("Money in"),
        _("Money out"),
        _("Added"),
        _("Statement reference"),
    ]
    return Sheet(_("Counterparties"), headers, rows, money_columns=(8, 9))


def region_sheet() -> Sheet:
    """Every region with its sales manager, so the counterparties download can be loaded into an empty database."""
    return Sheet(
        _("Regions"),
        [_("Region"), _("Sales manager"), _("Active")],
        [[region.name, region.manager, _yes_no(region.is_active)] for region in Region.objects.order_by("name")],
    )


def seed_sheets() -> list[Sheet]:
    """Everything the handbooks need on a new server: articles, categories, regions and counterparties."""
    return [*handbook_sheets(), region_sheet(), counterparty_sheet(Counterparty.objects.order_by("name"))]


# --- The project's handbook files ----------------------------------------------------
#
# The handbooks as Excel files kept with the code (cashbook/seed/), so a new server
# with an empty database can be filled from them: Handbook › Load from Excel ›
# "Load the project files", or ``manage.py load_handbooks``. They hold real
# counterparties (INNs, account numbers): they are served only to signed-in users
# with the page's permission, never as public static files, and the repository
# must stay private. ``manage.py export_handbooks --seed`` writes them again.

SEED_DIR = Path(__file__).resolve().parent / "seed"
SEED_FILES = {
    # kind: (file in SEED_DIR, download name); the handbook first: counterparties refer to its articles
    "handbook": ("handbook.xlsx", gettext_noop("Handbook")),
    "counterparties": ("counterparties.xlsx", gettext_noop("Counterparties")),
}


def seed_path(kind: str) -> Path:
    return SEED_DIR / SEED_FILES[kind][0]


def seed_sheets_by_kind() -> dict[str, list[Sheet]]:
    return {
        "handbook": handbook_sheets(),
        "counterparties": [counterparty_sheet(Counterparty.objects.order_by("name")), region_sheet()],
    }


def write_seed_files(directory: Path = SEED_DIR) -> list[Path]:
    """Write the project's handbook files from the database (in Russian, as the company uses them)."""
    directory.mkdir(parents=True, exist_ok=True)
    written = []
    with translation.override("ru"):
        for kind, sheets in seed_sheets_by_kind().items():
            path = directory / SEED_FILES[kind][0]
            path.write_bytes(xlsx_bytes(sheets))
            written.append(path)
    return written


def seed_response(kind: str) -> HttpResponse | None:
    """The project's file of ``kind`` as a download, or None if the project has none."""
    path = seed_path(kind)
    if not path.is_file():
        return None
    filename = f"{_(SEED_FILES[kind][1])} ({_('project file')}).xlsx"
    return attachment(path.read_bytes(), filename)
