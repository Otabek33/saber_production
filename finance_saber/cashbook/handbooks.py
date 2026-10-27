"""Load the handbooks (articles, categories, regions, counterparties) from Excel.

Reads the workbooks the "Download Excel" buttons produce (see exports.py), so a
new server with an empty database gets the same handbooks as the old one: export
them with ``manage.py export_handbooks`` (or the buttons), then upload the file
on the Handbook page or run ``manage.py load_handbooks``.

Any of the sheets may be present: "Статьи и категории" (one row per category),
"Статьи", "Регионы", "Контрагенты". Sheet names, column headers, article types
and yes / no are recognised in Russian and English, whichever language the file
was downloaded in. Records are matched by name (a counterparty by its statement
reference), so loading the same file again changes nothing. If any row has an
error, nothing is saved.

The finance workbook (Situation.xlsx) works too: its "Handbook" sheet lists the
bank counterparties with their article (columns H:I, region in J) and again with
article and region (Y:AA). From it, counterparties are created, and those without
an article or region in the app get the workbook's; ones already classified in
the app keep their classification.
"""

from collections import Counter
from dataclasses import dataclass, field

from django.core.exceptions import ValidationError
from django.db import transaction
from django.utils import translation
from django.utils.translation import gettext as _
from django.utils.translation import gettext_noop

from .bank import find_sheet
from .models import Article, ArticleKind, CashDesk, Category, Counterparty, Region

LANGUAGES = ("ru", "en")

# Sheets, and their columns (header -> field), as exports.py writes them.
FLAT = gettext_noop("Articles and categories")
ARTICLES = gettext_noop("Articles")
REGIONS = gettext_noop("Regions")
COUNTERPARTIES = gettext_noop("Counterparties")
COLUMNS = {
    FLAT: {
        gettext_noop("Type"): "kind",
        gettext_noop("Article"): "article",
        gettext_noop("Category"): "category",
        gettext_noop("Cash desks"): "desks",
        gettext_noop("Counts in cash flow"): "in_cash_flow",
        gettext_noop("Article active"): "is_active",
        gettext_noop("Category active"): "category_active",
    },
    ARTICLES: {
        gettext_noop("Type"): "kind",
        gettext_noop("Article"): "article",
        gettext_noop("Cash desks"): "desks",
        gettext_noop("Position"): "position",
        gettext_noop("Counts in cash flow"): "in_cash_flow",
        gettext_noop("Active"): "is_active",
    },
    REGIONS: {
        gettext_noop("Region"): "name",
        gettext_noop("Sales manager"): "manager",
        gettext_noop("Active"): "is_active",
    },
    COUNTERPARTIES: {
        gettext_noop("Name"): "name",
        gettext_noop("INN"): "inn",
        gettext_noop("Account number"): "account_number",
        gettext_noop("Article"): "article",
        gettext_noop("Region"): "region",
        gettext_noop("Sales manager"): "manager",
        gettext_noop("Statement reference"): "reference",
    },
}
# The finance workbook's counterparty handbook: (reference, article, region) column positions, 0-based.
FINANCE_WORKBOOK_SHEET = "Handbook"
FINANCE_WORKBOOK_BLOCKS = ((7, 8, 9), (24, 25, 26))  # H:J and Y:AA

REQUIRED = {
    FLAT: ("kind", "article"),
    ARTICLES: ("kind", "article"),
    REGIONS: ("name",),
    COUNTERPARTIES: ("reference",),
}


def texts(message: str) -> set[str]:
    """``message`` in every language, lower case, for matching what a file says."""
    found = {message.casefold()}
    for language in LANGUAGES:
        with translation.override(language):
            found.add(_(message).casefold())
    return found


def clean(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return " ".join(str(value).replace("\xa0", " ").split())


class HandbookError(Exception):
    """The file can't be read as a handbook workbook."""


@dataclass
class HandbookImport:
    sheets: list[str] = field(default_factory=list)  # the sheets found, as named in the file
    created: Counter = field(default_factory=Counter)  # by record type: "articles", "categories", ...
    updated: Counter = field(default_factory=Counter)
    unchanged: Counter = field(default_factory=Counter)
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    saved: bool = False

    def rows(self):
        """(label, created, updated, unchanged) per record type, for the result table."""
        models = {"articles": Article, "categories": Category, "regions": Region, "counterparties": Counterparty}
        labels = {key: str(model._meta.verbose_name_plural).capitalize() for key, model in models.items()}
        return [
            (label, self.created[key], self.updated[key], self.unchanged[key])
            for key, label in labels.items()
            if self.created[key] or self.updated[key] or self.unchanged[key]
        ]


class _Rollback(Exception):
    pass


class HandbookLoader:
    """Upserts handbook records from rows of the sheets above. Call inside a transaction."""

    def __init__(self, report: HandbookImport):
        self.report = report
        self.kinds = {kind.value: kind for kind in ArticleKind}
        for language in LANGUAGES:
            with translation.override(language):
                self.kinds.update({str(kind.label).casefold(): kind for kind in ArticleKind})
        self.yes = texts(gettext_noop("yes")) | {"true", "1", "да"}
        self.no = texts(gettext_noop("no")) | {"false", "0", "нет"}
        self.all_desks = texts(gettext_noop("All cash desks"))
        self.needs_article = texts(gettext_noop("Needs an article"))
        self.desks = {desk.name.casefold(): desk for desk in CashDesk.objects.all()}
        self.articles = {article.name.casefold(): article for article in Article.objects.prefetch_related("desks")}
        self.regions = {region.name.casefold(): region for region in Region.objects.all()}
        self.handled = set()  # articles already created or updated by this load
        self.unknown_articles = Counter()  # finance workbook: article names not in the handbook

    # --- values ---

    def error(self, sheet: str, row: int, message: str):
        self.report.errors.append(_("Sheet «%(sheet)s», row %(row)s: %(message)s") % {
            "sheet": sheet, "row": row, "message": message,
        })  # fmt: skip

    def flag(self, value, default=True):
        text = clean(value).casefold()
        if not text:
            return default
        if text in self.yes:
            return True
        if text in self.no:
            return False
        raise ValueError(_("“%(value)s” is not yes or no.") % {"value": clean(value)})

    def kind(self, value) -> str:
        kind = self.kinds.get(clean(value).casefold())
        if kind is None:
            choices = ", ".join(str(choice.label) for choice in ArticleKind)
            raise ValueError(_("unknown type “%(value)s” (expected: %(choices)s).") % {
                "value": clean(value), "choices": choices,
            })  # fmt: skip
        return kind

    def desk_list(self, value) -> list[CashDesk]:
        text = clean(value)
        if not text or text.casefold() in self.all_desks:
            return []
        desks = []
        for name in (part.strip() for part in text.split(",")):
            desk = self.desks.get(name.casefold())
            if desk is None:
                raise ValueError(_("unknown cash desk “%(name)s”.") % {"name": name})
            desks.append(desk)
        return desks

    def save(self, kind: str, instance, values: dict, created: bool):
        changed = created or any(getattr(instance, name) != value for name, value in values.items())
        for name, value in values.items():
            setattr(instance, name, value)
        if changed:
            instance.full_clean()
            instance.save()
        bucket = self.report.created if created else self.report.updated if changed else self.report.unchanged
        bucket[kind] += 1
        return changed

    # --- records ---

    def article(self, row: dict, with_position: bool) -> Article:
        name = clean(row["article"])
        article = self.articles.get(name.casefold())
        created = article is None
        if created:
            article = Article(name=name)
        values = {
            "kind": self.kind(row["kind"]),
            "in_cash_flow": self.flag(row.get("in_cash_flow")),
            "is_active": self.flag(row.get("is_active")),
        }
        if with_position and clean(row.get("position")).isdigit():
            values["position"] = int(clean(row["position"]))
        elif created:
            values["position"] = 999
        desks = self.desk_list(row.get("desks"))
        changed = self.save("articles", article, values, created)
        if {desk.pk for desk in article.desks.all()} != {desk.pk for desk in desks}:
            article.desks.set(desks)
            if not changed:  # counted as unchanged a moment ago
                self.report.unchanged["articles"] -= 1
                self.report.updated["articles"] += 1
        self.articles[name.casefold()] = article
        self.handled.add(name.casefold())
        return article

    def article_and_category(self, row: dict):
        """A row of "Статьи и категории": its article (the first time it appears in this load), then the category."""
        name = clean(row["article"]).casefold()
        article = self.articles[name] if name in self.handled else self.article(row, with_position=False)
        self.category(article, row)

    def category(self, article: Article, row: dict):
        name = clean(row.get("category"))
        if not name:
            return
        category = article.categories.filter(name=name).first()
        created = category is None
        if created:
            category = Category(article=article, name=name)
        self.save("categories", category, {"is_active": self.flag(row.get("category_active"))}, created)

    def region(self, row: dict) -> Region:
        name = clean(row["name"])
        region = self.regions.get(name.casefold())
        created = region is None
        if created:
            region = Region(name=name)
        values = {"manager": clean(row.get("manager")), "is_active": self.flag(row.get("is_active"))}
        self.save("regions", region, values, created)
        self.regions[name.casefold()] = region
        return region

    def counterparty(self, row: dict):
        reference = clean(row["reference"])
        parsed = Counterparty.parse_reference(reference)
        article_name = clean(row.get("article"))
        article = None
        if article_name and article_name.casefold() not in self.needs_article:
            article = self.articles.get(article_name.casefold())
            if article is None:
                raise ValueError(
                    _("unknown article “%(name)s”: add it to the handbook first.") % {"name": article_name}
                )
        region_name = clean(row.get("region"))
        region = None
        if region_name:
            region = self.regions.get(region_name.casefold()) or self.region(
                {"name": region_name, "manager": row.get("manager")}
            )
        counterparty = Counterparty.objects.filter(reference=reference).first()
        created = counterparty is None
        if created:
            counterparty = Counterparty(reference=reference)
        values = {
            "name": clean(row.get("name")) or parsed["name"],
            "inn": clean(row.get("inn")) or parsed["inn"],
            "account_number": clean(row.get("account_number")) or parsed["account_number"],
            "article": article,
            "region": region,
        }
        self.save("counterparties", counterparty, values, created)

    def counterparty_from_finance_workbook(self, row: dict):
        """A counterparty of Situation.xlsx's "Handbook" sheet; an article or region set in the app is kept."""
        reference = clean(row["reference"])
        counterparty = Counterparty.objects.filter(reference=reference).first()
        created = counterparty is None
        if created:
            counterparty = Counterparty(reference=reference, **Counterparty.parse_reference(reference))
        values = {}
        article_name = clean(row.get("article"))
        if article_name and counterparty.article_id is None:
            article = self.articles.get(article_name.casefold())
            if article is None:
                self.unknown_articles[article_name] += 1
            else:
                values["article"] = article
        region_name = clean(row.get("region"))
        if region_name and counterparty.region_id is None:
            values["region"] = self.regions.get(region_name.casefold()) or self.region({"name": region_name})
        self.save("counterparties", counterparty, values, created)


def _finance_workbook_table(sheet) -> list[tuple[int, dict]]:
    """Counterparties of the finance workbook's "Handbook" sheet, one per reference, with its first row number."""
    found = {}
    for number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        values = list(values) + [None] * 27
        for reference, article, region in FINANCE_WORKBOOK_BLOCKS:
            text = clean(values[reference])
            if "/" not in text:  # not a statement reference ("account/INN/name"), e.g. a total
                continue
            row = found.setdefault(text, (number, {"reference": text, "article": "", "region": ""}))[1]
            row["article"] = row["article"] or clean(values[article])
            region_text = clean(values[region])
            if region_text and region_text != "0":
                row["region"] = row["region"] or region_text
    return list(found.values())


def table(sheet, headers: dict[str, str], required=()) -> list[tuple[int, dict]]:
    """The sheet's rows as {field: value}, with their row numbers.

    Columns are found by their header (``headers``: header -> field), in any of
    the LANGUAGES; a missing ``required`` column is a HandbookError.
    """
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    wanted = {text: name for header, name in headers.items() for text in texts(header)}
    columns = {}
    for position, header in enumerate(rows[0]):
        name = wanted.get(clean(header).casefold())
        if name and name not in columns:
            columns[name] = position
    missing = [name for name in required if name not in columns]
    if missing:
        labels = {name: header for header, name in headers.items()}
        raise HandbookError(
            _("Sheet «%(sheet)s» has no column %(columns)s.") % {
                "sheet": sheet.title, "columns": ", ".join(f"«{_(labels[name])}»" for name in missing),
            }
        )  # fmt: skip
    table = []
    for number, values in enumerate(rows[1:], start=2):
        row = {name: values[position] if position < len(values) else None for name, position in columns.items()}
        if any(clean(value) for value in row.values()):
            table.append((number, row))
    return table


def load_handbooks(file, dry_run: bool = False) -> HandbookImport:
    """Load whichever handbook sheets the workbook has; nothing is saved on errors or a dry run."""
    import openpyxl

    try:
        workbook = openpyxl.load_workbook(file, data_only=True, read_only=True)
    except Exception as error:  # zip / xml errors: not a readable .xlsx
        raise HandbookError(_("The file is not a readable .xlsx workbook.")) from error
    report = HandbookImport()
    try:
        sheets = {}
        for title in (REGIONS, ARTICLES, FLAT, COUNTERPARTIES):
            for text in texts(title):
                found = find_sheet(workbook.sheetnames, text)
                if found:
                    sheets[title] = workbook[found]
                    report.sheets.append(found)
                    break
        tables = {title: table(sheet, COLUMNS[title], REQUIRED[title]) for title, sheet in sheets.items()}
        finance_sheet = find_sheet(workbook.sheetnames, FINANCE_WORKBOOK_SHEET)
        if finance_sheet:
            sheets[FINANCE_WORKBOOK_SHEET] = workbook[finance_sheet]
            tables[FINANCE_WORKBOOK_SHEET] = _finance_workbook_table(workbook[finance_sheet])
            report.sheets.append(finance_sheet)
        if not sheets:
            expected = [f"«{_(title)}»" for title in (FLAT, ARTICLES, REGIONS, COUNTERPARTIES)]
            names = ", ".join([*expected, f"«{FINANCE_WORKBOOK_SHEET}»"])
            raise HandbookError(_("No handbook sheets found. Expected any of: %(names)s.") % {"names": names})
        report.sheets.sort(key=workbook.sheetnames.index)  # as the file has them
    finally:
        workbook.close()

    try:
        with transaction.atomic():
            loader = HandbookLoader(report)
            steps = {
                REGIONS: loader.region,
                ARTICLES: lambda row: loader.article(row, with_position=True),
                FLAT: loader.article_and_category,
                COUNTERPARTIES: loader.counterparty,
                FINANCE_WORKBOOK_SHEET: loader.counterparty_from_finance_workbook,
            }
            # The order records depend on each other: regions and articles before counterparties.
            for title in (REGIONS, ARTICLES, FLAT, COUNTERPARTIES, FINANCE_WORKBOOK_SHEET):
                for number, row in tables.get(title, []):
                    try:
                        steps[title](row)
                    except (ValueError, ValidationError) as error:
                        message = "; ".join(error.messages) if isinstance(error, ValidationError) else str(error)
                        loader.error(sheets[title].title, number, message)
            for name, count in loader.unknown_articles.most_common():
                report.warnings.append(
                    _("Article «%(name)s» is not in the handbook: %(count)s counterparties left without an article.")
                    % {"name": name, "count": count}
                )
            if report.errors or dry_run:
                raise _Rollback
            report.saved = True
    except _Rollback:
        pass
    return report


def load_files(paths, dry_run: bool = False) -> HandbookImport:
    """Load several handbook files in order, all or nothing (e.g. the project's handbook, then counterparties).

    They are loaded in one transaction, so a dry run checks the later files
    against what the earlier ones would add (counterparties against new articles).
    """
    combined = HandbookImport()
    try:
        with transaction.atomic():
            for path in paths:
                report = load_handbooks(path)
                combined.sheets += report.sheets
                for counter in ("created", "updated", "unchanged"):
                    getattr(combined, counter).update(getattr(report, counter))
                combined.errors += report.errors
                combined.warnings += report.warnings
                if report.errors:
                    break
            if combined.errors or dry_run:
                raise _Rollback
            combined.saved = True
    except _Rollback:
        pass
    return combined
